"""Genera los documentos de la entrega del DMS en docs/dms/entrega:
  - Entrega-DMS-CADUS.docx (REM, análisis, diseño, casos de uso, implementación en SharePoint, trazabilidad y SCRUM)
  - casos-de-uso/CU-*.docx (un documento de casos de uso por tipo de documento: 4 documentos, el del acta y el del expediente con 2 CU)
Uso: python scripts/dms/build_doc_dms.py   (después, scripts/dms/exportar_pdf.ps1 para los PDF)"""
import csv
import json
import os
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DMS = os.path.join(RAIZ, "docs", "dms")
IMG = os.path.join(DMS, "img")
CAP = os.path.join(DMS, "capturas")
SAL = os.path.join(DMS, "entrega")
os.makedirs(os.path.join(SAL, "casos-de-uso"), exist_ok=True)
D = json.load(open(os.path.join(RAIZ, "scripts", "remus", "dms_data.json"), encoding="utf-8"))
AZUL = RGBColor(0x1F, 0x3A, 0x5F)


# ------------------------------------------------------------------ utilidades de Word
def nuevo():
    Num.fig = 0
    Num.tab = 0
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for n, tam in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        h = doc.styles[n]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = AZUL
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    return doc


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)


class Num:
    fig = 0
    tab = 0


def par(doc, texto, negrita=False, cursiva=False, tam=None, alinear=None):
    p = doc.add_paragraph()
    texto = texto.replace("`", "")
    partes = re.split(r"(\*\*.+?\*\*)", texto)
    for x in partes:
        if x.startswith("**") and x.endswith("**"):
            r = p.add_run(x[2:-2]); r.bold = True
        else:
            r = p.add_run(x); r.bold = negrita
        r.italic = cursiva
        if tam:
            r.font.size = Pt(tam)
    if alinear:
        p.alignment = alinear
    p.paragraph_format.space_after = Pt(6)
    return p


def vineta(doc, texto):
    p = doc.add_paragraph(style="List Bullet")
    partes = re.split(r"(\*\*.+?\*\*)", texto)
    for x in partes:
        if x.startswith("**") and x.endswith("**"):
            p.add_run(x[2:-2]).bold = True
        else:
            p.add_run(x)
    return p


def tabla(doc, cabecera, filas, anchos=None, titulo=None, tam=9):
    if titulo:
        Num.tab += 1
        p = doc.add_paragraph()
        r = p.add_run(f"Tabla {Num.tab}. {titulo}"); r.bold = True; r.font.size = Pt(9.5)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(3)
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        celda = t.rows[0].cells[i]
        celda.text = ""
        r = celda.paragraphs[0].add_run(c); r.bold = True; r.font.size = Pt(tam); r.font.color.rgb = RGBColor(255, 255, 255)
        sombrear(celda, "1F3A5F")
    for f in filas:
        cel = t.add_row().cells
        for i, v in enumerate(f):
            cel[i].text = ""
            r = cel[i].paragraphs[0].add_run(str(v).replace("`", "").replace("**", "")); r.font.size = Pt(tam)
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figura(doc, ruta, pie, ancho=16.5):
    Num.fig += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Cm(ancho))
    p.paragraph_format.keep_with_next = True
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {Num.fig}. {pie}"); r.italic = True; r.font.size = Pt(9)


# ------------------------------------------------------------------ burndown
def burndown():
    ruta = os.path.join(DMS, "burndown", "sprint-01.csv")
    filas = list(csv.DictReader(open(ruta, encoding="utf-8")))
    x = [f["fecha"][5:] for f in filas]
    ideal = [float(f["ideal_sp"]) for f in filas]
    real = [(i, float(f["real_sp"])) for i, f in enumerate(filas) if f["real_sp"] != ""]
    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=130)
    ax.plot(range(len(x)), ideal, "--", color="#888", label="Ideal")
    ax.plot([i for i, _ in real], [v for _, v in real], "-o", color="#1f3a5f", lw=2, label="Real")
    ax.set_xticks(range(len(x)))
    ax.set_xticklabels(x, rotation=45, fontsize=8)
    ax.set_ylabel("Puntos de historia pendientes")
    ax.set_title("Burndown del Sprint 1 del DMS (26/10 – 08/11/2026)", fontsize=11, color="#1f3a5f", fontweight="bold")
    ax.grid(alpha=.3)
    ax.legend()
    sal = os.path.join(IMG, "burndown-dms.png")
    fig.savefig(sal, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return sal


# ------------------------------------------------------------------ datos de los 4 tipos de documento
CAPS = {
    "01": ("01-sitios-y-bibliotecas.jpg", "Los tres sitios de grupo del DMS y sus bibliotecas en SharePoint de la US (sección «Compilar», «De su propiedad»).",
           "Se crearon los sitios «CADUS Documentación pública», «CADUS Órganos de gobierno» y «CADUS Atención al estudiante», cada uno con su biblioteca de documentos (Normativa, Actas y Votaciones, Expedientes). Cada sitio de grupo trae los grupos Propietarios, Integrantes y Visitantes que pide el Lab 4 y tiene 5 miembros: Pablo como propietario y los otros cuatro integrantes del equipo. Responde a RF-D01 y RF-D06 (DMS-HU-01)."),
    "02": ("02-normativa.jpg", "Biblioteca Normativa con sus columnas de metadatos.",
           "La biblioteca guarda la normativa del CADUS. Además del nombre, cada documento tiene las columnas Tipo de documento (obligatoria, con valor Normativa por defecto), Fecha, Curso académico y Estado (Vigente o Derogada). En la captura se ve un reglamento derogado junto a dos vigentes. Responde a RF-D02 y RN-D06 (DMS-HU-02)."),
    "03": ("03-normativa-vigente.jpg", "Vista plana «Vigente» de la biblioteca Normativa.",
           "La vista «Vigente» recorre la biblioteca sin carpetas y filtra por Estado = Vigente, así que el reglamento derogado desaparece. Es el filtro por metadatos de CU-D03 y de RF-D05 (DMS-HU-03 y DMS-HU-05)."),
    "04": ("04-actas.jpg", "Biblioteca Actas vista por la Ejecutiva (propietario): borrador y acta aprobada.",
           "La biblioteca Actas tiene activada la aprobación de contenido y la visibilidad de borradores solo para autor y aprobadores. La Ejecutiva ve a la vez el acta aprobada del pleno del 12/10 y el borrador de la asamblea del 26/10, con sus columnas Tipo de documento, Fecha, Curso académico y Asamblea. Responde a RF-D03 y RN-D01 (DMS-HU-04)."),
    "05": ("05-actas-aprobadas.jpg", "Vista plana «Aprobadas» de la biblioteca Actas.",
           "La vista «Aprobadas» filtra por estado de aprobación y solo muestra el acta aprobada, que es lo que debe ver un delegado en CU-D02. El borrador de la asamblea del 26/10 no aparece. Responde a RF-D03, RF-D05 y RN-D01."),
    "06": ("06-votaciones.jpg", "Biblioteca Votaciones con las columnas propias del registro de votación.",
           "La biblioteca Votaciones añade las columnas Número de votación, Tipo de voto (Nominal o Secreto) y Resultado. Los dos registros de ejemplo son de la misma asamblea: la votación 1 nominal y la 2 secreta. Los documentos no contienen números de DNI y, en la secreta, solo el recuento total. Responde a CU-D06, RN-D04 y RN-D05 (DMS-HU-07). En la biblioteca «Plantillas» del mismo sitio hay una plantilla de acta y otra de registro de votación con una fila por cada una de las 27 delegaciones (RF-D09)."),
    "07": ("07-permisos-votaciones.jpg", "Permisos exclusivos de la biblioteca Votaciones.",
           "La biblioteca Votaciones rompe la herencia de permisos del sitio: «Esta biblioteca tiene permisos exclusivos». Solo conserva Propietarios (control total) e Integrantes (editar); se ha quitado el grupo Visitantes, en el que irán los delegados, para que no vean los registros de votación. Responde a RF-D06 y RN-D04."),
    "08": ("08-expedientes.jpg", "Biblioteca Expedientes del sitio Atención al estudiante.",
           "La biblioteca guarda los expedientes de reclamaciones y mediaciones resueltos, con las columnas Tipo de documento, Fecha, Curso académico, Código (CAD-… o MED-…), Facultad, Estado y Fecha de resolución, más la fecha de caducidad calculada. Tiene también permisos exclusivos, sin el grupo Visitantes (RN-D02). Responde a CU-D04 y RF-D01 (DMS-HU-06)."),
    "09": ("09-expedientes-caducados.jpg", "Vista plana «Caducados» de la biblioteca Expedientes.",
           "La columna Fecha de caducidad es calculada (fecha de resolución más 2 años). La vista «Caducados» filtra los expedientes cuya caducidad es anterior a hoy: aparece el MED-2024-0002, resuelto el 10/05/2024 y caducado desde el 10/05/2026. Es lo que la Ejecutiva revisa cada mes para retirar los expedientes (CU-D05, RF-D07, RF-D08, RN-D03)."),
}

TIPOS = [
    {"archivo": "CU-Acta", "nombre": "Acta de pleno o asamblea", "cus": ["CU-D01", "CU-D02"], "sitio": "CADUS Órganos de gobierno", "biblioteca": "Actas",
     "desc": "Documento que recoge lo tratado y acordado en un pleno o una asamblea. Se redacta a partir de una plantilla, queda en borrador hasta que la Ejecutiva la aprueba y solo entonces la ven los delegados.",
     "meta": [["Tipo de documento", "Elección (Acta)", "Sí"], ["Fecha", "Fecha", "Sí"], ["Curso académico", "Elección", "No"], ["Asamblea", "Texto", "No"], ["Estado de aprobación", "Nativo de SharePoint (borrador, aprobada)", "Automático"], ["Resumen", "Generado automáticamente (RF-D10, deseable)", "No"]],
     "permisos": "Propietarios (cuenta del CADUS): leer y escribir. Visitantes (delegados): leer, solo las actas aprobadas. Los borradores solo los ven el autor y los aprobadores.",
     "caps": ["04", "05"], "reglas": ["RN-D01"], "ejemplo": "Acta del pleno del 12-10-2026 (ejemplo).docx (aprobada) y Acta de la asamblea del 26-10-2026 (ejemplo, borrador).docx"},
    {"archivo": "CU-Normativa", "nombre": "Normativa y estatutos", "cus": ["CU-D03"], "sitio": "CADUS Documentación pública", "biblioteca": "Normativa",
     "desc": "Normas, reglamentos y estatutos del CADUS. Solo se muestran las vigentes en la vista principal; las derogadas se conservan como referencia.",
     "meta": [["Tipo de documento", "Elección (Normativa)", "Sí"], ["Fecha", "Fecha", "Sí"], ["Curso académico", "Elección", "No"], ["Estado", "Elección (Vigente, Derogada)", "Sí"]],
     "permisos": "Propietarios: leer y escribir. Visitantes (estudiantes y delegados): solo lectura.",
     "caps": ["02", "03"], "reglas": [], "ejemplo": "Estatutos del CADUS (ejemplo).docx, Reglamento de votaciones (ejemplo).docx y un reglamento derogado"},
    {"archivo": "CU-Expediente", "nombre": "Expediente de reclamación o mediación", "cus": ["CU-D04", "CU-D05"], "sitio": "CADUS Atención al estudiante", "biblioteca": "Expedientes",
     "desc": "Expediente de una reclamación o mediación ya resuelta. Se conserva 2 años desde su resolución como prueba y luego se retira. Solo lo ve la Ejecutiva.",
     "meta": [["Tipo de documento", "Elección (Expediente)", "Sí"], ["Fecha", "Fecha", "Sí"], ["Curso académico", "Elección", "No"], ["Código", "Texto (CAD-… o MED-…)", "No"], ["Facultad", "Elección", "No"], ["Estado", "Elección (Resuelto)", "Sí"], ["Fecha de resolución", "Fecha (obligatoria: de ella depende la caducidad)", "Sí"], ["Fecha de caducidad", "Calculada: resolución + 2 años", "Automático"]],
     "permisos": "Permisos exclusivos: Propietarios (control total) e Integrantes (editar, solo el equipo con datos de ejemplo). Sin el grupo Visitantes.",
     "caps": ["08", "09"], "reglas": ["RN-D02", "RN-D03"], "ejemplo": "CAD-2026-0001 (ejemplo).docx y MED-2024-0002 (ejemplo, caducado).docx"},
    {"archivo": "CU-Votacion", "nombre": "Registro de votación", "cus": ["CU-D06"], "sitio": "CADUS Órganos de gobierno", "biblioteca": "Votaciones",
     "desc": "Resultado de una votación presencial de asamblea, con un voto por delegación (27 delegaciones). No guarda el DNI y, si la votación es secreta, solo el recuento total.",
     "meta": [["Tipo de documento", "Elección (Votación)", "Sí"], ["Fecha", "Fecha", "Sí"], ["Curso académico", "Elección", "No"], ["Asamblea", "Texto", "No"], ["Número de votación", "Número", "No"], ["Tipo de voto", "Elección (Nominal, Secreto)", "Sí"], ["Resultado", "Elección (Aprobada, Rechazada)", "No"]],
     "permisos": "Permisos exclusivos: Propietarios (control total) e Integrantes (editar, el equipo del proyecto). Sin el grupo Visitantes: los delegados no ven los registros.",
     "caps": ["06", "07"], "reglas": ["RN-D04", "RN-D05"], "ejemplo": "Votación 1 (nominal, ejemplo).docx y Votación 2 (secreta, ejemplo).docx"},
]
CU = {u["id"]: u for u in D["usecases"]}
ACT = {a["id"]: a["name"] for a in D["actors"]}
RF = {r["id"]: r for r in D["functional"]}


def bloque_cu(doc, cu, nivel=2):
    doc.add_heading(f"{cu['id']} {cu['name']}", nivel)
    actores = sorted({ACT[s["actor"]] for s in cu["steps"] if s["who"] == "actor"})
    if cu["id"] == "CU-D03":
        actores = ["Delegado", "Estudiante"]
    tabla(doc, ["Campo", "Contenido"], [["Actor", ", ".join(actores)], ["Cuándo", cu["trigger"][0].upper() + cu["trigger"][1:]], ["Precondición", cu["pre"]], ["Postcondición", cu["post"]]], [3.5, 13.5], tam=9)
    par(doc, "**Pasos**")
    for i, s in enumerate(cu["steps"], 1):
        quien = ACT[s["actor"]] if s["who"] == "actor" else "El sistema"
        par(doc, f"{i}. {quien} {s['text']}.")
    trazas = next((t[1] for t in D["traces"] if t[0] == cu["id"]), [])
    par(doc, "**Requisitos relacionados:** " + ", ".join(trazas))


def doc_cu(tipo):
    doc = nuevo()
    t = doc.add_heading(f"Casos de uso — {tipo['nombre']}", 0)
    par(doc, "DMS del CADUS · Lab 4 · Universidad de Sevilla (GSI) · versión 1.0", cursiva=True, tam=9.5)
    doc.add_heading("1. Documento", 1)
    par(doc, tipo["desc"])
    tabla(doc, ["Dato", "Valor"], [["Sitio de SharePoint", tipo["sitio"]], ["Biblioteca", tipo["biblioteca"]], ["Casos de uso", ", ".join(tipo["cus"])], ["Reglas de negocio", ", ".join(tipo["reglas"]) or "—"], ["Documentos de ejemplo", tipo["ejemplo"]]], [4, 13], tam=9)
    doc.add_heading("2. Metadatos", 1)
    tabla(doc, ["Metadato", "Tipo de columna", "Obligatorio"], tipo["meta"], [5, 9, 3], tam=9)
    doc.add_heading("3. Permisos", 1)
    par(doc, tipo["permisos"])
    doc.add_heading("4. Casos de uso", 1)
    for c in tipo["cus"]:
        bloque_cu(doc, CU[c], 2)
    doc.add_heading("5. Implementación en SharePoint", 1)
    for k in tipo["caps"]:
        arch, pie, texto = CAPS[k]
        par(doc, texto)
        figura(doc, os.path.join(CAP, arch), pie, 16)
    sal = os.path.join(SAL, "casos-de-uso", tipo["archivo"] + ".docx")
    doc.save(sal)
    return sal


# ------------------------------------------------------------------ documento de la entrega
def backlog():
    filas = []
    for l in open(os.path.join(DMS, "backlog", "product-backlog.md"), encoding="utf-8"):
        c = [x.strip() for x in l.strip().strip("|").split("|")]
        if len(c) == 8 and c[0].startswith("DMS-HU"):
            filas.append([c[0], c[1], c[2], c[4], c[5], c[7]])
    return filas


def entrega():
    doc = nuevo()
    doc.add_heading("DMS del CADUS — Entrega del sistema de gestión documental", 0)
    par(doc, "Gestión de Sistemas de Información (GSI) · Universidad de Sevilla · Lab 4 (DMS)", cursiva=True)
    par(doc, "Equipo: Pablo Ruiz Vidal (Product Owner), Joaquín Luna Canela (Scrum Master), Miguel Ángel Camacho Martín, Álvaro Martínez Ocaña y Marcos-Paban Delgado Rodríguez. Cliente: CADUS (contacto: Miguel Ángel Camacho Martín).")
    par(doc, "Fecha de entrega: 08/11/2026 · Versión 1.0 (estado a 30/10/2026)")

    doc.add_heading("1. Resumen", 1)
    par(doc, "El DMS guarda y organiza los documentos internos del CADUS —actas, normativa, expedientes de reclamaciones y registros de votación— en **SharePoint de Microsoft 365 de la Universidad de Sevilla**. Es independiente del portal web (CMS): tiene su propio REM, análisis, diseño, implementación y seguimiento SCRUM. Los requisitos salen de la entrevista con el CADUS del 26/10/2026.")
    tabla(doc, ["Contenido pedido para el DMS", "Dónde está"], [
        ["Requisitos (REM)", "Sección 2 y proyecto REMUS `DMS.rem` (más la hoja `carga-remus.md`)"],
        ["Documentos de análisis (UML)", "Sección 3 (UML de análisis y casos de uso)"],
        ["Documentos de diseño (UML)", "Sección 4 (UML de diseño y del análisis al diseño)"],
        ["Documentos de casos de uso (3-4 documentos, uno con más de 1 CU)", "Sección 5 y carpeta `casos-de-uso`: 4 documentos (acta y expediente con 2 CU cada uno)"],
        ["Documentos de implementación (SharePoint)", "Sección 6: 9 capturas explicadas una a una"],
        ["Seguimiento SCRUM del DMS", "Sección 8: product backlog, sprint backlog y burndown"],
        ["Entrevista con el cliente", "Sección 2.9 y anexo (sección 10); también como documento aparte"]], [7, 10])

    doc.add_heading("2. Requisitos (REM)", 1)
    par(doc, "El REM se mantiene en REMUS (`DMS.rem`, versión 1.0) a partir de `scripts/remus/dms_data.json`. Cada requisito lleva un código que se repite en el UML, en los casos de uso, en las capturas y en la trazabilidad.")
    doc.add_heading("2.1 Objetivos", 2)
    tabla(doc, ["Código", "Objetivo", "El sistema deberá"], [[o["id"], o["name"], o["description"]] for o in D["objectives"]], [2, 4, 11], tam=9)
    doc.add_heading("2.2 Actores", 2)
    tabla(doc, ["Código", "Actor", "Representa a"], [[a["id"], a["name"], a["role"]] for a in D["actors"]], [2, 4, 11], tam=9)
    doc.add_heading("2.3 Requisitos de información", 2)
    tabla(doc, ["Código", "Concepto", "Datos"], [[r["id"], r["name"], "; ".join(r["data"])] for r in D["information"]], [2, 3.5, 11.5], tam=8.5)
    doc.add_heading("2.4 Requisitos de interfaz y funcionales", 2)
    tabla(doc, ["Código", "Nombre", "El sistema deberá", "Historia"], [[r["id"], r["name"], r["description"], r["comments"]] for r in D["ui"] + D["functional"]], [1.8, 3.4, 9.2, 2.6], tam=8.5)
    doc.add_heading("2.5 Requisitos no funcionales", 2)
    tabla(doc, ["Código", "Nombre", "El sistema deberá"], [[r["id"], r["name"], r["description"]] for r in D["nonfunctional"]], [2, 3.5, 11.5], tam=9)
    doc.add_heading("2.6 Reglas de negocio", 2)
    tabla(doc, ["Código", "Nombre", "Regla"], [[r["id"], r["name"], r["description"]] for r in D["rules"]], [2, 3.5, 11.5], tam=9)
    doc.add_heading("2.7 Restricciones del entorno", 2)
    tabla(doc, ["Código", "Nombre", "Descripción"], [[r["id"], r["name"], r["description"]] for r in D["environment"]], [2, 4, 11], tam=9)
    doc.add_heading("2.8 Casos de uso", 2)
    tabla(doc, ["Caso de uso", "Quién", "Cuándo", "Documento"], [
        ["CU-D01 Redactar, aprobar y archivar un acta", "Ejecutiva", "Después de cada pleno o asamblea", "Acta"],
        ["CU-D02 Consultar las actas aprobadas", "Delegado", "Antes de cada asamblea", "Acta"],
        ["CU-D03 Consultar la normativa vigente", "Estudiante o delegado", "Cuando tienen una duda", "Normativa"],
        ["CU-D04 Archivar un expediente resuelto", "Ejecutiva", "Al resolver una reclamación o mediación", "Expediente"],
        ["CU-D05 Retirar los expedientes caducados", "Ejecutiva", "Una vez al mes", "Expediente"],
        ["CU-D06 Registrar el resultado de una votación", "Ejecutiva", "En cada asamblea con votación", "Registro de votación"]], [6.5, 3, 4.5, 3], tam=9)
    doc.add_heading("2.9 Origen de los requisitos: la entrevista", 2)
    par(doc, "Las respuestas del cliente que moldean el DMS (entrevista del 26/10/2026): las votaciones son presenciales, solo votan los delegados con un voto por delegación y se comprueba la identidad con el DNI; no se publican resultados en el portal; los tickets y mediaciones se conservan 2 años desde que se resuelven; todas las consultas deben tratarse con total privacidad; y todo pasa por la Ejecutiva, que usa la cuenta única del CADUS. El texto completo está en `docs/cliente/entrevista-cliente.md`.")

    doc.add_heading("3. Análisis", 1)
    par(doc, "El UML de análisis recoge los conceptos del Lab 4: documento (con sus cuatro tipos), carpeta, rol, permiso, búsqueda y metadato. Las etiquetas en rojo son los códigos de requisito.")
    figura(doc, os.path.join(IMG, "uml-dms-analisis.png"), "UML de análisis del DMS con etiquetas de trazabilidad.", 17)
    par(doc, "El diagrama de casos de uso relaciona cada actor con los casos de uso que ejecuta. Cada tipo de documento tiene al menos un caso de uso y el acta y el expediente tienen dos.")
    figura(doc, os.path.join(IMG, "casos-de-uso-dms.png"), "Casos de uso del DMS.", 16)

    doc.add_heading("4. Diseño", 1)
    par(doc, "En el diseño cada concepto del análisis pasa a un elemento de SharePoint: los documentos tienen columnas de metadatos, las carpetas son sitios y bibliotecas, los roles son grupos de permisos y la búsqueda es una búsqueda genérica más vistas planas con filtros.")
    figura(doc, os.path.join(IMG, "uml-dms-diseno.png"), "UML de diseño del DMS en SharePoint.", 17)
    tabla(doc, ["Concepto del análisis", "Elemento del diseño en SharePoint"], [
        ["Documento", "Documento de una biblioteca, con columnas de metadatos"],
        ["Carpeta", "Sitio de grupo y biblioteca de documentos"],
        ["Rol (Ejecutiva, Delegado, Estudiante)", "Grupo de permisos: Propietarios, Integrantes y Visitantes"],
        ["Permiso", "Permiso del grupo sobre el sitio o la biblioteca (permisos exclusivos en Votaciones y Expedientes)"],
        ["Búsqueda", "Búsqueda de SharePoint por nombre y contenido"],
        ["Metadato", "Columna de la biblioteca (elección, fecha, texto, número o calculada)"],
        ["Filtro", "Vista plana de la biblioteca con filtro por metadatos (Vigente, Aprobadas, Caducados)"]], [6.5, 10.5], tam=9, titulo="Del análisis al diseño")
    doc.add_heading("4.1 Metadatos y vistas por biblioteca", 2)
    tabla(doc, ["Sitio", "Biblioteca", "Columnas propias", "Vista plana"], [
        ["Documentación pública", "Normativa", "Tipo de documento, Fecha, Curso académico, Estado", "Vigente (Estado = Vigente)"],
        ["Órganos de gobierno", "Actas", "Tipo de documento, Fecha, Curso académico, Asamblea; aprobación de contenido", "Aprobadas (estado de aprobación = Aprobado)"],
        ["Órganos de gobierno", "Votaciones", "Tipo de documento, Fecha, Curso académico, Asamblea, Número de votación, Tipo de voto, Resultado", "Todas las votaciones"],
        ["Atención al estudiante", "Expedientes", "Tipo de documento, Fecha, Curso académico, Código, Facultad, Estado, Fecha de resolución, Fecha de caducidad (calculada)", "Caducados (caducidad anterior a hoy)"]], [3.3, 2.7, 7.3, 3.7], tam=8.5, titulo="Bibliotecas, columnas y vistas")

    doc.add_heading("5. Documentos de casos de uso", 1)
    par(doc, "Se entregan **4 documentos de casos de uso**, uno por tipo de documento, y dos de ellos (acta y expediente) tienen más de un caso de uso. Cada documento describe el tipo de documento, sus metadatos, sus permisos, sus casos de uso paso a paso y la captura de SharePoint donde se ve.")
    tabla(doc, ["Documento de casos de uso", "Casos de uso", "Biblioteca"], [[f"{t['archivo']}.docx — {t['nombre']}", ", ".join(t["cus"]), f"{t['sitio']} · {t['biblioteca']}"] for t in TIPOS], [7.5, 3.5, 6], tam=9)
    for t in TIPOS:
        doc.add_heading(t["nombre"], 2)
        par(doc, t["desc"])
        for c in t["cus"]:
            u = CU[c]
            par(doc, f"**{c} {u['name']}.** " + " ".join(f"{i}. {(ACT[s['actor']] if s['who'] == 'actor' else 'El sistema')} {s['text']}." for i, s in enumerate(u["steps"], 1)))

    doc.add_heading("6. Implementación en SharePoint", 1)
    par(doc, "Capturas del entorno real (`uses0.sharepoint.com`), con documentos de ejemplo ficticios. Cada una indica los requisitos y la historia a la que responde.")
    for k in sorted(CAPS):
        arch, pie, texto = CAPS[k]
        doc.add_heading(f"{int(k)}. {pie.rstrip('.')}", 3)
        par(doc, texto)
        figura(doc, os.path.join(CAP, arch), pie, 16)

    doc.add_heading("7. Trazabilidad", 1)
    tabla(doc, ["Caso de uso", "Requisitos", "Historia", "Elemento de SharePoint", "Captura", "Estado"], [
        ["CU-D01", "RF-D01, D02, D03, D09, RN-D01", "DMS-HU-01, 02, 04", "Biblioteca Actas con aprobación de contenido; plantilla de acta", "4", "Configurado y observado"],
        ["CU-D02", "RF-D03 a D06, RN-D01", "DMS-HU-03, 04", "Vista plana Aprobadas", "5", "Configurado; falta probar con cuenta de delegado"],
        ["CU-D03", "RF-D04 a D06", "DMS-HU-03, 05", "Vista plana Vigente de Normativa", "3", "Configurado y observado"],
        ["CU-D04", "RF-D01, D02, D07, RN-D02", "DMS-HU-06", "Biblioteca Expedientes (permisos exclusivos)", "8", "Configurado y observado"],
        ["CU-D05", "RF-D07, D08, RN-D03, RNF-D04", "DMS-HU-06", "Vista plana Caducados y columna calculada", "9", "Configurado y observado; falta probar la retirada"],
        ["CU-D06", "RF-D01, D02, D09, RN-D04, RN-D05", "DMS-HU-07", "Biblioteca Votaciones (permisos exclusivos) y plantilla de votación", "6, 7", "Configurado y observado"]], [1.6, 3.8, 2.4, 4.6, 1.2, 3.4], tam=8.5, titulo="Casos de uso, requisitos, historias y SharePoint")
    par(doc, "Los objetivos OBJ-D01 a OBJ-D04 y los requisitos de información IR-D01 a IR-D06 están enlazados con los requisitos y casos de uso en las trazas de `DMS.rem` (71 trazas).")

    doc.add_heading("8. Seguimiento SCRUM del DMS", 1)
    par(doc, "El DMS tiene su propio product backlog, sprint backlog y burndown, independientes de los del CMS. Sprint 1 del DMS: 26/10 – 08/11/2026 (24 SP).")
    doc.add_heading("8.1 Product backlog", 2)
    tabla(doc, ["ID", "Historia", "Requisitos", "SP", "Estado", "Responsable"], backlog(), [2, 6.5, 3, 1, 2.5, 2], tam=8)
    doc.add_heading("8.2 Sprint backlog", 2)
    tabla(doc, ["Integrante", "Historias del DMS", "SP"], [["Pablo", "DMS-HU-01", 5], ["Joaquín", "DMS-HU-02, DMS-HU-05", 7], ["Miguel Ángel", "DMS-HU-04, DMS-HU-07", 6], ["Álvaro", "DMS-HU-03", 3], ["Marcos", "DMS-HU-06", 3]], [4, 10, 2], tam=9)
    par(doc, "Tareas sin puntos: D-01 REM en REMUS (Pablo), D-02 y D-03 UML de análisis y diseño (Álvaro), D-04 capturas (todos), D-05 seguimiento SCRUM (Joaquín) y D-06 este documento (Pablo).")
    doc.add_heading("8.3 Burndown del sprint", 2)
    figura(doc, burndown(), "Burndown chart del Sprint 1 del DMS: puntos pendientes por día frente a la línea ideal.", 14)
    par(doc, "La configuración de las siete historias se hizo en bloque el 30/10 y el burndown no se actualizó los días anteriores: repite la lección de los sprints del CMS. Según la definición de hecho del sprint (caso de uso ejecutado por otra persona y permisos comprobados con cada rol), ninguna historia está todavía cerrada: seis están configuradas a la espera de la prueba cruzada y DMS-HU-04 espera además una cuenta de delegado. Por eso el burndown sigue en 24 SP; bajará cuando el equipo cierre las pruebas.")

    doc.add_heading("9. Pendiente y limitaciones", 1)
    for t in ("Probar los permisos con cuentas reales de delegado y de estudiante: ahora los grupos Visitantes están vacíos.",
              "Añadir la cuenta del CADUS (dcadus@us.es) como propietaria de los tres sitios.",
              "Prueba cruzada de cada caso de uso por otra persona del equipo (definición de hecho del sprint).",
              "RF-D10 (resumen automático del acta) es deseable y no se ha implementado.",
              "Los documentos de ejemplo son ficticios; los datos reales los aporta el CADUS.",
              "Estado de verificación: lo que está **configurado** (sitios, columnas, vistas y permisos), **observado** (capturas) y **probado** (por otra persona o con otra cuenta) no es lo mismo; hasta ahora solo hay configuración y observación con la cuenta de Pablo"):
        vineta(doc, t)
    anexo_entrevista(doc)
    sal = os.path.join(SAL, "Entrega-DMS-CADUS.docx")
    doc.save(sal)
    return sal


# ------------------------------------------------------------------ entrevista con el cliente
def leer_entrevista():
    """Lee docs/cliente/entrevista-cliente.md: cabecera, bloques de preguntas y tabla de consecuencias."""
    lineas = open(os.path.join(RAIZ, "docs", "cliente", "entrevista-cliente.md"), encoding="utf-8").read().splitlines()
    cab, bloques, tabla_c, cur, modo = [], [], [], None, "cab"
    i = 0
    while i < len(lineas):
        l = lineas[i].strip()
        if l.startswith("## Qué cambia"):
            modo = "tabla"
        elif l.startswith("## "):
            cur = {"titulo": re.sub(r"^## \d+\.\s*", "", l), "preguntas": []}
            bloques.append(cur)
            modo = "bloque"
        elif modo == "cab" and l and not l.startswith("#"):
            cab.append(l.replace("**", ""))
        elif modo == "bloque" and l.startswith("**") and l.endswith("**"):
            q = re.sub(r"^\*\*(\d+)\.\s*", r"\1. ", l)[:-2]
            resp = []
            i += 1
            while i < len(lineas) and lineas[i].strip() and not lineas[i].strip().startswith("**"):
                resp.append(lineas[i].strip())
                i += 1
            cur["preguntas"].append((q, " ".join(resp)))
            continue
        elif modo == "tabla" and l.startswith("|") and not l.startswith("|---"):
            tabla_c.append([c.strip().replace("`", "") for c in l.strip("|").split("|")])
        i += 1
    return cab, bloques, tabla_c


def cuerpo_entrevista(doc, nivel):
    cab, bloques, tabla_c = leer_entrevista()
    for c in cab:
        par(doc, c)
    for b in bloques:
        doc.add_heading(b["titulo"], nivel + 1)
        for q, r in b["preguntas"]:
            par(doc, "**" + q + "**")
            par(doc, r)
    if tabla_c:
        doc.add_heading("Qué cambia en el proyecto según las respuestas", nivel)
        tabla(doc, tabla_c[0], tabla_c[1:], [6.2, 8.3, 2.5], tam=8.5)


def entrevista_docx():
    doc = nuevo()
    doc.add_heading("Entrevista con el CADUS", 0)
    par(doc, "Gestión de Sistemas de Información (GSI) · Universidad de Sevilla · requisitos del CMS y del DMS", cursiva=True, tam=9.5)
    cuerpo_entrevista(doc, 1)
    sal = os.path.join(SAL, "Entrevista-CADUS.docx")
    doc.save(sal)
    return sal


def anexo_entrevista(doc):
    doc.add_heading("10. Anexo: entrevista con el CADUS", 1)
    par(doc, "Texto completo de la entrevista con el cliente en la que se basan los requisitos del DMS y los cambios del CMS. También se adjunta como documento aparte (Entrevista-CADUS).")
    cuerpo_entrevista(doc, 2)


if __name__ == "__main__":
    for t in TIPOS:
        print("generado", doc_cu(t))
    print("generado", entrevista_docx())
    print("generado", entrega())
