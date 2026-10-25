"""Genera el documento Word de la entrega. Uso: python scripts/entrega/build_doc.py
Requiere: python-docx. Antes, ejecutar diagramas.py para crear las imágenes."""
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
import tempfile

import numpy as np
from PIL import Image

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ENTREGA = os.path.join(RAIZ, "docs", "entrega")
IMG = os.path.join(ENTREGA, "img")
CAP = os.path.join(IMG, "capturas")
SALIDA = os.path.join(ENTREGA, "Entrega2-CMS-CADUS.docx")
FECHA_DATOS = "24/10/2026"

datos = json.load(open(os.path.join(RAIZ, "scripts", "remus", "cadus_data.json"), encoding="utf-8"))
plugins = {p["slug"]: p for p in json.load(open(os.path.join(ENTREGA, "datos", "plugins.json"), encoding="utf-8"))}
temas = {t["slug"]: t for t in json.load(open(os.path.join(ENTREGA, "datos", "themes.json"), encoding="utf-8"))}
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
ROJO = RGBColor(0xB3, 0x26, 0x1E)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
for nombre, tam in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
    h = doc.styles[nombre]
    h.font.name = "Calibri"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.font.size = Pt(tam)
    h.font.bold = True
    h.font.color.rgb = AZUL
    h.paragraph_format.space_before = Pt(14 if nombre != "Heading 1" else 18)
    h.paragraph_format.space_after = Pt(6)
cap = doc.styles["Caption"]
cap.font.size = Pt(9)
cap.font.italic = True
cap.font.bold = False
cap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

figura_n = [0]
tabla_n = [0]


# ------------------------------------------------------------------ utilidades
def campo(par, codigo, texto_previo=""):
    for tipo in ("begin", None, "separate", "texto", "end"):
        r = par.add_run()
        if tipo in ("begin", "separate", "end"):
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
            r._r.append(f)
        elif tipo is None:
            i = OxmlElement("w:instrText")
            i.set(qn("xml:space"), "preserve")
            i.text = codigo
            r._r.append(i)
        else:
            r.text = texto_previo


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def par(texto="", negrita=False, cursiva=False, color=None, tam=None, alin=None, estilo=None, despues=None):
    p = doc.add_paragraph(style=estilo)
    if texto:
        r = p.add_run(texto)
        r.bold, r.italic = negrita, cursiva
        if color:
            r.font.color.rgb = color
        if tam:
            r.font.size = Pt(tam)
    if alin:
        p.alignment = alin
    if despues is not None:
        p.paragraph_format.space_after = Pt(despues)
    return p


def rico(texto, estilo=None):
    """Párrafo con **negrita** y `código` en línea."""
    p = doc.add_paragraph(style=estilo)
    for trozo in re.split(r"(\*\*.+?\*\*|`.+?`)", texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            p.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`"):
            r = p.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            p.add_run(trozo)
    return p


def vineta(texto):
    return rico(texto, "List Bullet")


def tabla(cabecera, filas, anchos, tam=8.5, titulo=None, resaltar=None):
    if titulo:
        tabla_n[0] += 1
        c = doc.add_paragraph(style="Caption")
        c.add_run(f"Tabla {tabla_n[0]}. {titulo}")
        c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(cabecera):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(tam)
        r.font.color.rgb = RGBColor(255, 255, 255)
        sombrear(c, "1F3A5F")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for n, fila in enumerate(filas):
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v))
            r.font.size = Pt(tam)
            if resaltar and resaltar(fila):
                r.bold = True
            if n % 2 == 1:
                sombrear(cells[i], "F2F5F9")
    for row in t.rows:
        for i, w in enumerate(anchos):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figura(archivo, texto, ancho=16.5, carpeta=IMG, max_h=22.0, caja=None):
    """Inserta una imagen con su pie. Si es muy alta se reduce para que quepa en la página; caja=(y0, y1) recorta un trozo."""
    ruta = os.path.join(carpeta, archivo)
    if caja:
        base = Image.open(ruta)
        ruta = os.path.join(tempfile.gettempdir(), f"trozo_{caja[0]}_{archivo}")
        base.crop((0, caja[0], base.width, caja[1])).convert("RGB").save(ruta, quality=90)
    w, h = Image.open(ruta).size
    ancho = min(ancho, max_h * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(ruta, width=Cm(ancho))
    figura_n[0] += 1
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.add_run(f"Figura {figura_n[0]}. {texto}")
    return figura_n[0]


def cortes(archivo, carpeta, n):
    """Divide una captura muy alta en n trozos, cortando por filas lisas (entre secciones) cuando es posible."""
    a = np.asarray(Image.open(os.path.join(carpeta, archivo)).convert("RGB")).astype(int)
    alto = len(a)
    puntos = [0]
    for i in range(1, n):
        objetivo = alto * i // n
        corte = objetivo
        for d in range(0, 450, 3):
            ok = [y for y in (objetivo - d, objetivo + d) if 0 < y < alto and a[y].max() - a[y].min() < 4]
            if ok:
                corte = ok[0]
                break
        puntos.append(corte)
    puntos.append(alto)
    return list(zip(puntos[:-1], puntos[1:]))


def leer_seccion(ruta, titulo):
    """Líneas de la sección «## titulo» de un Markdown, hasta la siguiente sección."""
    dentro, salida = False, []
    for linea in open(ruta, encoding="utf-8").read().splitlines():
        if linea.startswith("## "):
            dentro = linea[3:].strip() == titulo
            continue
        if dentro and linea.strip():
            salida.append(linea)
    return salida


def md_linea(linea):
    """Un párrafo o una viñeta de Markdown con negrita y código; quita la cursiva y los enlaces."""
    linea = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", linea)
    linea = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", linea)
    if linea.startswith("- "):
        vineta(linea[2:])
    elif not linea.startswith(("|", ">")):
        rico(linea)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def millar(n):
    if not n:
        return "—"
    if n >= 1_000_000:
        return f"{n // 1_000_000} M+"
    return f"{n // 1000} k+"


def fecha(iso):
    a, m, d = iso.split("-")
    return f"{d}/{m}/{a}"


def estrellas(r):
    return f"{r / 20:.1f}/5".replace(".", ",")


def comparar(titulo, slugs, elegido, motivo, tema=False):
    fuente = temas if tema else plugins
    filas = []
    for s in slugs:
        p = fuente[s]
        nombre = re.sub(r"\s+[–—-]\s+.*$", "", p["nombre"]).strip()[:34]
        filas.append([nombre, p.get("version", ""), fecha(p["actualizado"]), millar(p.get("instal")), estrellas(p["rating"]), f'{p["votos"]:,}'.replace(",", "."),
                      "Seleccionado" if s == elegido else ""])
    tabla(["Plugin" if not tema else "Tema", "Versión", "Última actualización", "Instalaciones activas", "Valoración", "Nº valoraciones", "Decisión"], filas,
          [4.0, 1.7, 2.5, 2.4, 1.9, 2.0, 2.5], titulo=titulo, resaltar=lambda f: f[-1] == "Seleccionado")
    rico("**Justificación.** " + motivo)


def parsear_tablas(ruta):
    tablas, actual = [], None
    for linea in open(ruta, encoding="utf-8"):
        linea = linea.rstrip("\n")
        if linea.startswith("|"):
            celdas = [c.strip().strip("`") for c in linea.strip().strip("|").split("|")]
            if re.match(r"^:?-+:?$", celdas[0]):
                continue
            if actual is None:
                actual = []
                tablas.append(actual)
            actual.append(celdas)
        else:
            actual = None
    return tablas


# ------------------------------------------------------------------ portada
for _ in range(3):
    par()
par("Universidad de Sevilla · Gestión de Sistemas de Información (GSI)", color=RGBColor(0x55, 0x55, 0x55), alin=WD_ALIGN_PARAGRAPH.CENTER, tam=12)
par("Portal web del CADUS", negrita=True, color=AZUL, tam=30, alin=WD_ALIGN_PARAGRAPH.CENTER, despues=2)
par("Sistema 1 · CMS en WordPress", color=AZUL, tam=16, alin=WD_ALIGN_PARAGRAPH.CENTER)
par("Entrega 2 · Versión final · 25 de octubre de 2026", negrita=True, tam=13, alin=WD_ALIGN_PARAGRAPH.CENTER, despues=24)
tabla(["Integrante", "Rol SCRUM / Área"], [[s["name"], s["role"]] for s in datos["stakeholders"][1:]], [7.0, 10.0], tam=10, titulo="Equipo del proyecto")
par("Cliente: CADUS – Consejo de Alumnos de la Universidad de Sevilla · Pabellón de Uruguay, Av. de Chile, s/n, 41013 Sevilla · cadus@us.es", cursiva=True, tam=9.5, alin=WD_ALIGN_PARAGRAPH.CENTER)
salto()

par("Índice", negrita=True, color=AZUL, tam=18)
campo(par(), 'TOC \\o "1-2" \\h \\z \\u', "Haz clic derecho sobre el índice y elige «Actualizar campo» para generarlo.")
salto()

# ------------------------------------------------------------------ resumen
doc.add_heading("Resumen del estado de la entrega", 1)
par("Esta es la segunda entrega del CMS (versión final, 25 de octubre): incorpora el Sprint 2 completo y mantiene lo ya entregado el 18 de octubre. El portal está en funcionamiento sobre WordPress 7.1 con todas las historias de usuario del alcance (HU-01 a HU-13, HU-18 y HU-19). Lo que falta, y por qué, está en la sección 8.")
n = lambda k: len(datos[k])
tabla(["Apartado solicitado", "Estado", "Sección"], [
    ["Planificación: herramientas, product backlog, sprint backlog y burndown del Sprint 2", "Completo: Sprint 2 cerrado con 29 de 29 SP; seguimiento en GitHub Project", "2"],
    [f"REM revisado (versión 2.0): {n('objectives')} objetivos, {n('information')} de información, {n('ui')} de interfaz, {n('functional')} funcionales, {n('nonfunctional')} no funcionales, {len(datos['rules'])} reglas y {n('usecases')} casos de uso", "REM 2.0 generado y documentado (PDF del documento que genera REMUS: CADUS-REM-v2.pdf). Falta la entrevista con el CADUS para validarlo: sigue pendiente", "3"],
    ["Análisis: estructura web, diagrama de clases, mockups y capturas de menús y de la página principal", "Completo con 7 mockups de baja fidelidad y 3 temáticas visuales comparadas; los mockups están también en Figma (bocetos de baja fidelidad)", "4"],
    ["Diseño: comparativa de al menos tres plugins por tipo (con los obligatorios) y de al menos tres temas", "Completo: 11 comparativas de plugins por tipo y 5 temas", "5"],
    ["Implementación: capturas explicadas una a una", "Completa para el alcance salvo la papeleta de las votaciones (RF-12 parcial: se publican las convocatorias, no se puede votar); el contenido sigue siendo de ejemplo", "6"],
    ["Trazabilidad entre REM, análisis, diseño e implementación", "Completo", "7"],
    ["Calidad, incidencias y pendientes", "Revisión cruzada Claude–Codex, incidencias y decisiones abiertas", "8"],
], [9.5, 5.5, 2.0], titulo="Estado de cada apartado de la entrega")
par("Nota sobre el sprint: se completaron los 29 SP comprometidos, pero el trabajo se concentró otra vez en los dos últimos días (23 de los 29 puntos, el 79 %, se cerraron el 23 y el 24 de octubre) y el burndown no se actualizó cada día como se había acordado. La retrospectiva (sección 2.7) recoge las acciones para el Sprint 3.", cursiva=True)

# ------------------------------------------------------------------ 1. propuesta
doc.add_heading("1. Propuesta de temática y empresa (versión definitiva)", 1)
tabla(["Campo", "Valor"], [
    ["Organización cliente", "CADUS – Consejo de Alumnos de la Universidad de Sevilla"],
    ["Sector", "Educación superior / representación estudiantil"],
    ["Persona de contacto", "Delegación General del CADUS (cadus@us.es)"],
    ["Ubicación", "Pabellón de Uruguay, Av. de Chile, s/n, 41013 Sevilla"],
    ["Sistema", "CMS en WordPress (Sistema 1 del proyecto en grupo)"],
], [4.5, 12.5], titulo="Datos del cliente")
rico(datos["paragraphs"]["intro"])
doc.add_heading("Objetivos del sistema", 2)
for o in datos["objectives"]:
    vineta(f"**{o['id']} · {o['name']}:** {o['description']}.")

# ------------------------------------------------------------------ 2. planificación
doc.add_heading("2. Planificación", 1)
doc.add_heading("2.1 Equipo y organización", 2)
par("El equipo lo forman cinco integrantes con roles SCRUM y áreas funcionales diferenciadas. Cada miembro debe poder operar y explicar los tres sistemas, por lo que el reparto por áreas no implica aislamiento de conocimiento.")
tabla(["Integrante", "Rol SCRUM y área funcional"], [[s["name"], s["role"]] for s in datos["stakeholders"][1:]], [6.5, 10.5], titulo="Roles del equipo")

doc.add_heading("2.2 Comparativa y justificación de herramientas", 2)
doc.add_heading("2.2.1 Hosting de WordPress", 3)
par("Se comparan cuatro alternativas para alojar el CMS. La valoración es cualitativa; no se han contratado ni probado los servicios de pago.")
tabla(["Criterio", "Servidor local WampServer + Tailscale", "Hosting gratuito compartido", "Hosting compartido de pago", "WordPress.com"], [
    ["Coste", "0 €", "0 €", "Cuota periódica", "Plan gratuito limitado; los planes con plugins son de pago"],
    ["Control sobre PHP, MySQL y archivos", "Total", "Limitado", "Alto", "Bajo"],
    ["Instalación de plugins propios", "Sí", "Sí", "Sí", "Solo en planes de pago"],
    ["Disponibilidad", "Solo con el equipo anfitrión encendido", "Variable", "24 h / 7 d", "24 h / 7 d"],
    ["Acceso simultáneo del equipo", "Sí, por red privada Tailscale", "Sí, por internet", "Sí, por internet", "Sí, por internet"],
    ["Cumple lo que pide la guía del proyecto", "Sí: exige servidor local (XAMPP, WAMP o MAMP)", "Opcional", "Opcional", "Opcional"],
    ["Riesgo principal", "Dependencia de un solo PC", "Límites y rendimiento", "Coste", "Falta de control"],
], [3.4, 3.7, 3.0, 3.3, 3.6], tam=8, titulo="Comparativa de alojamiento para WordPress")
rico("**Elección: servidor local WampServer compartido mediante Tailscale.** La guía del proyecto exige un servidor local verificado con PHP y MySQL; Wamp cumple ese requisito sin coste, permite instalar cualquier plugin y deja que los cinco integrantes trabajen a la vez sobre la misma instancia de WordPress y la misma base de datos `bd_cadus`. La VPN privada de Tailscale evita abrir puertos del router y limita el acceso a los miembros invitados. El inconveniente asumido es la disponibilidad, que se mitiga avisando en Discord (#avisos-servidor) y manteniendo copias de la base de datos en el repositorio.")

doc.add_heading("2.2.2 Software de comunicación del equipo", 3)
tabla(["Criterio", "Discord", "Microsoft Teams", "WhatsApp", "Slack"], [
    ["Coste", "Gratuito", "Gratuito con cuenta universitaria o personal", "Gratuito", "Gratuito con límites"],
    ["Canales por tema", "Sí (categorías y canales)", "Sí (equipos y canales)", "No (grupos planos)", "Sí"],
    ["Voz y pantalla compartida para trabajar desde casa", "Sí, canales de voz permanentes", "Sí, reuniones", "Llamadas básicas", "Llamadas (limitadas en gratis)"],
    ["Integración con GitHub", "Sí, webhook de commits", "Sí, mediante aplicación", "No", "Sí"],
    ["Historial y búsqueda", "Ilimitado y buscable", "Sí", "Limitado a búsqueda simple", "Limitado en el plan gratuito"],
    ["Encaje en el proyecto", "Comunicación diaria del equipo", "Se usará como Sistema 3 del proyecto (Teams, Planner, Power Platform)", "Avisos urgentes", "—"],
], [3.4, 3.4, 3.8, 3.0, 3.4], tam=8, titulo="Comparativa de software de comunicación")
rico("**Elección: Discord.** Permite estructurar la comunicación en categorías (INFO, SCRUM, DESARROLLO y VOZ), mantiene salas de voz siempre abiertas para el trabajo a distancia y recibe los commits de GitHub. Teams se reserva para el Sistema 3 de la asignatura, en el que será objeto de implantación y no herramienta de trabajo. La estructura creada está descrita en `docs/arquitectura/entorno-de-trabajo.md`.")

doc.add_heading("2.2.3 Software para la gestión de SCRUM", 3)
tabla(["Criterio", "GitHub Projects", "Jira", "Trello", "Microsoft Planner", "Taiga"], [
    ["Coste", "Gratuito", "Gratuito para equipos pequeños", "Gratuito con límites", "Incluido en Microsoft 365", "Gratuito para proyectos pequeños"],
    ["Integración con el repositorio", "Nativa: incidencias y pull requests", "Mediante integración", "Mediante extensión", "No", "Mediante integración"],
    ["Puntos de historia y sprints", "Campos personalizados e iteraciones", "Sí", "Con extensiones", "No", "Sí"],
    ["Burndown", "No nativo: se calcula desde el backlog", "Nativo", "Con extensiones", "No", "Nativo"],
    ["Plantillas SCRUM", "Plantillas de incidencias y de tablero", "Sí", "Sí", "Básicas", "Sí"],
    ["Curva de aprendizaje", "Baja (ya se usa GitHub)", "Media-alta", "Baja", "Baja", "Media"],
], [3.0, 3.0, 2.6, 2.6, 3.0, 2.8], tam=8, titulo="Comparativa de herramientas de gestión SCRUM")
rico("**Elección: GitHub Projects**, con tablero, incidencias etiquetadas como historia de usuario (existen plantillas en `.github/ISSUE_TEMPLATE`) y los artefactos SCRUM versionados en el repositorio. Es la opción de menor fricción: el equipo ya trabaja en GitHub, que la guía define como fuente de la verdad, y cada historia puede enlazarse con sus commits y pull requests. Su carencia es el burndown nativo; se compensa generando el gráfico a partir de los puntos de historia del backlog (`docs/scrum/burndown/sprint-01.csv`, generado por `scripts/entrega/diagramas.py`). Si el equipo prefiere un burndown automático, Taiga sería la alternativa SCRUM puro.")

doc.add_heading("2.2.4 Sistema de control de versiones", 3)
tabla(["Criterio", "Git + GitHub", "Git + GitLab", "Git + Bitbucket", "Subversion"], [
    ["Modelo", "Distribuido", "Distribuido", "Distribuido", "Centralizado"],
    ["Coste para el equipo", "Gratuito", "Gratuito", "Gratuito", "Gratuito (requiere servidor)"],
    ["Revisión de cambios (pull / merge request)", "Sí", "Sí", "Sí", "No nativa"],
    ["Incidencias y tableros", "Sí (Projects)", "Sí", "Sí (con Jira)", "No"],
    ["Exigido por la guía del proyecto", "Sí: «fuente de la verdad»", "No", "No", "No"],
], [4.5, 3.5, 3.0, 3.0, 3.0], tam=8, titulo="Comparativa de control de versiones")
rico("**Elección: Git con GitHub.** La guía exige un repositorio GitHub con la estructura `docs/`, `scrum/`, `cms-wordpress/`, `dms-sharepoint/` y la carpeta del sistema 3, un `README.md` detallado y plantillas SCRUM. El flujo de ramas (`main`, `develop`, `feature/*`) y las plantillas de pull request están documentados en el README del repositorio.")

doc.add_heading("2.3 Product backlog", 2)
par("El product backlog es el mismo de todo el proyecto CMS. Se ordena por prioridad y agrupa tareas técnicas (PB) e historias de usuario (HU), cada una enlazada con los requisitos del REM. Estimación en puntos de historia (Fibonacci).")
tb = parsear_tablas(os.path.join(RAIZ, "docs", "scrum", "backlog", "product-backlog.md"))
infra = [[c[0], c[1], c[4], c[5], c[6]] for c in tb[0][1:]]
tabla(["ID", "Tarea técnica", "SP", "Estado", "Sprint"], infra, [1.6, 10.0, 1.1, 2.0, 2.3], titulo="Product backlog: tareas técnicas y de infraestructura")
hu = [[c[0], c[1], c[2], c[4], c[6]] for c in tb[1][1:]]
tabla(["ID", "Historia de usuario", "Requisitos", "SP", "Sprint"], hu, [1.5, 9.1, 3.0, 1.0, 2.4], tam=8, titulo="Product backlog: historias de usuario")
par(f"Total del product backlog: {sum(int(c[4]) for c in tb[0][1:]) + sum(int(c[4]) for c in tb[1][1:])} puntos de historia.", cursiva=True)

doc.add_heading("2.4 Sprint 1 (28/09 – 18/10/2026): resumen", 2)
rico("El Sprint 1 se entregó el 18/10/2026 y su sprint backlog y su burndown figuran en el documento de la entrega 1. **Se completaron los 70 SP comprometidos** (plataforma base, identidad visual, contenidos, directorio de delegaciones y buzón). Sus datos están en `docs/scrum/sprints/sprint-01.md` y `docs/scrum/burndown/sprint-01.csv`. Este documento recoge el sprint backlog y el burndown del sprint de esta entrega, el Sprint 2.")

doc.add_heading("2.5 Sprint backlog del Sprint 2 (19/10 – 25/10/2026)", 2)
rico("**Objetivo del Sprint 2:** completar el portal con las historias de participación (tickets avanzados, calendario, votaciones, mediación y panel de administración), sustituir el contenido de ejemplo por contenido real validado con el cliente y entregar la documentación final. **Alcance comprometido: 29 SP**, más las tareas de documentación `E2-xx`, que no llevan puntos. Contacto con el cliente: Miguel Ángel Camacho, que forma parte del CADUS.")
t2 = parsear_tablas(os.path.join(RAIZ, "docs", "scrum", "sprints", "sprint-02.md"))
personas = ["Pablo", "Joaquín", "Miguel Ángel", "Álvaro", "Marcos"]
estado2 = {"E2-01": "En curso", "E2-02": "Pendiente", "E2-03": "Hecho", "E2-04": "Pendiente", "E2-05": "Hecho", "E2-06": "Hecho", "E2-07": "Hecho",
           "E2-08": "Pendiente", "E2-09": "Pendiente", "E2-10": "Hecho", "E2-11": "Pendiente", "E2-12": "Parcial", "E2-13": "Hecho",
           "E2-14": "Hecho", "E2-15": "Hecho", "E2-16": "Hecho", "HU-19c": "Hecho"}
extra = [["Pablo", "E2-17", "Exportar el REM a PDF desde REMUS (la vista previa de REMUS sale en blanco; PDF obtenido del HTML que genera REMUS)", "—", "Hecho"],
         ["Pablo", "E2-18", "Fusionar las pull requests #32 y #33 de Álvaro (pruebas por rol y su documentación)", "—", "Hecho"],
         ["Miguel Ángel", "E2-19", "Crear 2 o 3 eventos reales en el calendario para la demostración", "—", "Pendiente"],
         ["Pablo", "E2-20", "Revisión cruzada final con Codex y corrección de los hallazgos", "—", "Hecho"]]
filas2 = []
for quien, t in zip(personas, [x for x in t2 if x[0][:2] == ["ID", "Tarea concreta"]]):
    for c in t[1:]:
        corto = re.sub(r"\s+", " ", c[1])
        corto = corto if len(corto) < 150 else corto[:147].rsplit(" ", 1)[0] + "…"
        filas2.append([quien, c[0], corto.replace("`", "").replace("**", ""), c[2], estado2.get(c[0], "Hecho")])
filas2 += extra  # tareas E2-17 a E2-20, añadidas durante el sprint
tabla(["Responsable", "ID", "Tarea", "SP", "Estado"], filas2, [2.4, 1.4, 9.2, 0.9, 3.1], tam=8, titulo="Sprint backlog del Sprint 2")
sp2 = sum(int(f[3]) for f in filas2 if f[3].isdigit())
par(f"Resultado: {sp2} de {sp2} SP completados (100 %). De las 20 tareas `E2-xx` de documentación y entrega, doce están hechas (E2-05 con la salvedad de que el burndown no se actualizó a diario, E2-10 con bocetos de baja fidelidad en Figma y E2-17 con el PDF obtenido del HTML que genera REMUS), una parcial, una en curso y seis pendientes por depender de datos del cliente o del cierre de la entrega (sección 8). El reparto final de puntos fue de 5 a 7 SP por persona.".replace("`", ""), cursiva=True)

doc.add_heading("2.6 Burndown chart del Sprint 2", 2)
figura("burndown-sprint2.png", "Burndown chart del Sprint 2 (alcance 29 SP). Ideal frente a real.", 13.5)
rico("El 19/10 se cerraron HU-19a y HU-09b (6 SP); el 23/10 HU-19b, HU-10 y HU-12 (11 SP) y el 24/10 HU-09a, HU-11 y HU-19c (12 SP), quedando 0 SP pendientes. **Lección:** se repitió lo del Sprint 1: 23 de los 29 puntos (el 79 %) se cerraron en los dos últimos días. Además, el burndown no se actualizó cada día como se había acordado: se reconstruyó al final con las fechas reales de los commits y de las pull requests. Datos en `docs/scrum/burndown/sprint-02.csv`; la gráfica y los datos se regeneran con `scripts/entrega/diagramas.py`.")

doc.add_heading("2.7 Sprint Review y retrospectiva del Sprint 2", 2)
rico("**Sprint Review (24/10/2026).** Se revisó el incremento sobre el servidor compartido: la portada (40 aniversario, accesos, noticias con filtros, eventos, «Nuestro trabajo», galería, buscador y pasos del buzón), la participación (calendario, convocatorias de votación y mediación), la atención al estudiante (buzón con código de seguimiento, panel de tickets con filtros, preguntas frecuentes y aviso de privacidad) y la administración (menús por rol, panel de inicio y cuentas del CADUS y de los delegados). Queda por validar con el cliente, a través de Miguel Ángel, los textos reales, los objetivos de «Nuestro trabajo» (proceden de la web oficial y son de 2024) y la lista de cuentas de las delegaciones.")
tabla(["Historia", "SP", "Quién", "Cierre", "Qué se entregó"], [
    ["HU-19a · HU-09b", "6", "Pablo", "19/10", "Menús de administración por rol y privacidad de los tickets (RN-03)"],
    ["HU-19b", "3", "Miguel Ángel", "23/10", "Panel de inicio con accesos según el rol"],
    ["HU-10", "3", "Miguel Ángel", "23/10", "Calendario de asambleas y eventos (The Events Calendar)"],
    ["HU-12", "5", "Álvaro", "23/10", "Solicitud de mediación privada con la protección RN-03"],
    ["HU-09a", "5", "Marcos", "24/10", "Filtros de tickets, aviso al estudiante sin contenido, FAQ y aviso de privacidad"],
    ["HU-11", "5", "Joaquín", "24/10", "Convocatorias de votación vinculadas a una asamblea (RN-05); la emisión del voto pasa a HU-21 (Sprint 3)"],
    ["HU-19c", "2", "Álvaro", "24/10", "Plan y ejecución de pruebas por rol (PR #32, fusionada el 25/10)"],
], [2.6, 0.9, 2.6, 1.5, 9.4], tam=8, titulo="Resultado del Sprint 2 por historia")
rico("**Qué ha ido bien.** La arquitectura por módulos del plugin permitió trabajar en paralelo (calendario, mediación, votaciones y buzón se integraron sin tocarse entre sí). La revisión cruzada entre Claude y Codex encontró fallos reales antes de entregar y la batería de pruebas de privacidad se amplió con cada módulo. El GitHub Project fue la fuente única del backlog.")
rico("**Qué ha ido mal.** Se repitió la concentración del trabajo al final y el burndown no se actualizó a diario. Una rama (la de mediación) se hizo sobre una copia antigua del plugin y habría eliminado la portada, los menús por rol y el diseño al fusionarse: se evitó integrando solo su módulo. Varios entregables (mockups en Figma, entrevista con el cliente y contenido real) no estaban listos al cerrar el sprint.")
par("Acciones para el Sprint 3:")
for t in ("Actualizar el tablero y el burndown cada día (cinco minutos en la daily) y no cerrar historias «de golpe».",
          "Trabajar siempre desde `git pull` de `main` y abrir una pull request pequeña por historia; no copiar carpetas de otra rama.",
          "Pedir al cliente los contenidos reales al inicio del sprint, no al final.",
          "Mantener la batería de pruebas de privacidad como requisito de cierre de cualquier historia que toque datos de estudiantes."):
    vineta(t)

doc.add_heading("2.8 Seguimiento en GitHub Project", 2)
rico("El backlog se gestiona en el proyecto público **«CADUS · Scrum»** de GitHub (`github.com/users/pablitoruiizz/projects/1`), vinculado al repositorio. Tiene dos vistas guardadas: **Product backlog** (tabla agrupada por sprint) y **Sprint backlog** (tablero del sprint actual). Cada ítem lleva los campos SP, Requisitos (códigos del REM), Prioridad, Responsable, Sprint (iteración) y Cierre (fecha real, de la que sale el burndown). Contiene las historias y tareas técnicas, las 20 tareas de documentación del Sprint 2 y las historias del Sprint 3 ya asignadas a su iteración. Los cinco integrantes tienen acceso de escritura.")

# ------------------------------------------------------------------ 3. REM
doc.add_heading("3. Requisitos (REM)", 1)
par("Los requisitos se han recogido en REMUS (fichero CADUS.rem, versión 2.0 del 24/10/2026) con la estructura de cinco dimensiones: información, interfaz, funcionales, no funcionales y reglas de negocio. Cada requisito lleva su código de trazabilidad, que se repite en el análisis, el diseño y las capturas. La versión 2.0 recoge las decisiones del Sprint 2: tickets con el módulo propio del buzón, mediación con referencia MED-AAAA-NNNNN, votaciones vinculadas a una asamblea, una cuenta única del CADUS (dcadus@us.es) que usan todos los miembros de la Ejecutiva y desde la que se da acceso a los delegados, y una cuenta por delegado con su correo de alumno (@alum.us.es), privacidad ampliada a las mediaciones, requisitos de interfaz del calendario, la mediación y el panel por rol, y dos casos de uso nuevos. El documento generado por REMUS se entrega aparte (CADUS-REM-v2.pdf) junto al fichero CADUS.rem; las tablas siguientes son un resumen de su contenido. Una diferencia conocida: el código de ticket (CAD-AAAA-NNNNN) y la referencia de mediación (MED-AAAA-NNNNN) llevan en el portal un sufijo aleatorio de 32 caracteres hexadecimales que el REM aún no recoge; se incorporará al validarlo con el cliente. La entrevista con el CADUS prevista para confirmar los requisitos no se ha realizado: el REM queda marcado como pendiente de validar con el cliente, que incluye el plazo de conservación de los datos, la lista de cuentas reales y el logotipo oficial. La importancia indicada es la asignada en REMUS.")
imp = {2: "Vital", 3: "Importante", 4: "Quedaría bien"}
doc.add_heading("3.1 Requisitos de información", 2)
tabla(["Código", "Nombre", "Se almacenará información correspondiente a…", "Datos específicos"],
      [[r["id"], r["name"], r["concept"], "; ".join(r["data"])] for r in datos["information"]], [1.5, 3.0, 5.0, 7.5], tam=8, titulo="Requisitos de información")
doc.add_heading("3.2 Actores", 2)
tabla(["Actor", "Representa a…"], [[a["name"], a["role"]] for a in datos["actors"]], [4.0, 13.0], titulo="Actores del sistema")
doc.add_heading("3.3 Requisitos de interfaz de usuario", 2)
tabla(["Código", "Nombre", "El sistema deberá…", "Historias"], [[r["id"], r["name"], r["description"].replace("\n\n- ", "; ").replace("\n", " "), r["comments"]] for r in datos["ui"]],
      [1.5, 3.3, 10.2, 2.0], tam=8, titulo="Requisitos de interfaz de usuario")
doc.add_heading("3.4 Requisitos funcionales", 2)
tabla(["Código", "Nombre", "El sistema deberá…", "Historias", "Importancia"], [[r["id"], r["name"], r["description"], r["comments"], imp[r["importance"]]] for r in datos["functional"]],
      [1.4, 3.5, 7.2, 2.0, 2.2], tam=8, titulo="Requisitos funcionales")
doc.add_heading("Casos de uso", 3)
for u in datos["usecases"]:
    rico(f"**{u['id']} · {u['name']}** — Evento disparador: {u['trigger']}. Precondición: {u['pre']}. Postcondición: {u['post']}.")
    for i, s in enumerate(u["steps"], 1):
        quien = next(a["name"] for a in datos["actors"] if a["id"] == s["actor"]) if s["who"] == "actor" else "El sistema"
        par(f"{i}. {quien} {s['text']}.", despues=1)
    par()
doc.add_heading("3.5 Requisitos no funcionales", 2)
tabla(["Código", "Nombre", "El sistema deberá…", "Importancia"], [[r["id"], r["name"], r["description"], imp[r["importance"]]] for r in datos["nonfunctional"]], [1.6, 3.6, 9.8, 2.0], tam=8.5, titulo="Requisitos no funcionales")
doc.add_heading("3.6 Reglas de negocio", 2)
tabla(["Código", "Nombre", "Regla"], [[r["id"], r["name"], r["description"]] for r in datos["rules"]], [1.6, 4.0, 11.4], titulo="Reglas de negocio")
doc.add_heading("3.7 Restricciones del entorno de trabajo", 2)
tabla(["Código", "Nombre", "Descripción"], [[r["id"], r["name"], r["description"]] for r in datos["environment"]], [1.6, 3.8, 11.6], tam=8, titulo="Restricciones del proyecto (Tailscale, GitHub y Discord)")

# ------------------------------------------------------------------ 4. análisis
doc.add_heading("4. Análisis", 1)
doc.add_heading("4.1 Estructura del sitio web", 2)
par("El portal se organiza en seis secciones de primer nivel colgando de la portada. Cada rama indica los requisitos que implementa.")
figura("uml-analisis-estructura-web.png", "Estructura de menús y páginas del portal, con las etiquetas de trazabilidad (IR, RF, UI).", 16)
doc.add_heading("4.2 Diagrama de clases", 2)
par("El modelo de información se deriva de los requisitos de información IR-03 a IR-09. Las entidades en azul tienen requisito propio; las verdes se han deducido de los requisitos funcionales.")
figura("uml-analisis-clases.png", "Diagrama de clases del análisis con las etiquetas de trazabilidad de cada entidad.", 17)
doc.add_heading("4.3 Mockups", 2)
par("Se muestran siete prototipos de baja fidelidad: los cuatro del Sprint 1 (portada, directorio de delegaciones, buzón de consultas y ficha de acta) y tres nuevos del Sprint 2 (calendario con votaciones, solicitud de mediación y panel de administración por rol). Cada uno indica los requisitos que cubre. Los siete mockups y las tres temáticas de la sección 4.4 están también en Figma, como capas editables y con los códigos de requisito en cada pantalla: https://www.figma.com/design/IZ4OfXGSS6T8WLToDq8f8s/CADUS---Mockups-entrega-2. Son los mismos bocetos de baja fidelidad; la maqueta definitiva de cada pantalla es el portal implementado, que se muestra en la sección 6.")
figura("mockup-01-portada.png", "Mockup de la portada. UI-01 · UI-02 · IR-01.", 13.5)
par("Portada: banner institucional con el lema, comunicados recientes por categoría, accesos destacados al directorio de delegaciones y al buzón, y una franja con las próximas asambleas, eventos y votaciones. El pie lleva los datos de contacto. Navegación: menú principal de siete entradas.")
figura("mockup-02-directorio.png", "Mockup del directorio de delegaciones y su ficha. UI-03 · RF-06 · RF-07.", 13.5)
par("Directorio: filtro por facultad o escuela y delegaciones agrupadas por centro. Al pulsar una delegación se abre su ficha con delegado, correo, teléfono, horario de atención y enlaces; solo el delegado de esa delegación puede editarla (RN-04).")
figura("mockup-03-buzon.png", "Mockup del buzón de consultas y del seguimiento de ticket. UI-04 · RF-08 · RF-09 · RN-02 · RN-03.", 13.5)
par("Buzón: formulario con correo @alum.us.es, asunto, tipo y descripción; a la derecha, la consulta del estado con el código de seguimiento (abierto, en proceso o cerrado) y el recuadro de privacidad (RN-03). Estados: envío correcto con código, error de dominio de correo (RN-02) y consulta sin coincidencia.")
figura("mockup-04-acta.png", "Mockup de la ficha de acta o normativa. UI-05 · RF-03 · RF-04 · RF-05 · RN-01.", 13.5)
par("Ficha de documento: vista previa del PDF, datos (tipo, fecha, pleno, estado y categoría) y botón de descarga. El recuadro recuerda el flujo de aprobación borrador → pendiente → aprobada (RN-01).")
figura("mockup-05-calendario.png", "Mockup del calendario de asambleas y eventos y de la convocatoria de votación. UI-06 · RF-11 · RF-12 · RN-05.", 13.5)
par("Calendario: buscador, vistas de lista, mes y día, y los próximos eventos con su fecha. A la derecha, la convocatoria de votación vinculada a una asamblea con su apertura y cierre; solo se puede votar con la votación abierta. Estados: sin eventos programados y votación próxima, abierta o cerrada.")
figura("mockup-06-mediacion.png", "Mockup de la solicitud de mediación y su confirmación. UI-07 · RF-13 · RN-02 · RN-03.", 13.5)
par("Mediación: formulario con nombre, correo @alum.us.es, facultad o escuela, descripción del conflicto y consentimiento RGPD. Tras enviarlo se muestra solo la referencia, sin repetir los datos; el recuadro explica que la solicitud es privada (RN-03).")
figura("mockup-07-panel.png", "Mockup del escritorio de administración y de los menús según el rol. UI-08 · RF-14 · RF-15.", 13.5)
par("Escritorio: menú lateral con las secciones del CADUS, tarjetas con lo pendiente (tickets, mediaciones, actas y delegaciones) y un recuadro que resume qué menús ve cada rol: la Ejecutiva todos, el delegado solo su ficha y el estudiante ninguno.")

doc.add_heading("4.4 Tres temáticas visuales para la comparativa", 2)
par("Para elegir la identidad visual del portal se diseñaron tres temáticas distintas de la portada, cada una asociada a un tema de WordPress candidato. Los bocetos usan la misma estructura (cabecera, banner, accesos, noticias y pie) para que la comparación sea justa; solo cambian la paleta, el estilo y la composición.")
TEM = json.load(open(os.path.join(ENTREGA, "datos", "tematicas.json"), encoding="utf-8"))
figura("tematica-a.png", "Temática A · Institucional CADUS (tema candidato: Blocksy). UI-01 · UI-02.", 13)
figura("tematica-b.png", "Temática B · Claro y sobrio (tema candidato: Astra). UI-01 · UI-02.", 13)
figura("tematica-c.png", "Temática C · Oscuro contemporáneo (tema candidato: Kadence). UI-01 · UI-02.", 13)
par("Métricas de la comparación: (1) valoración y número de valoraciones del tema candidato en wordpress.org; (2) etiqueta «Accessibility ready»; (3) contraste del texto principal sobre el fondo, del texto del banner, del botón y de las tarjetas, calculado con la fórmula de WCAG 2.1 (el mínimo del nivel AA es 4,5:1); (4) fidelidad a la identidad visual del CADUS (paleta, logotipo y tipografía de cadus.us.es).")
filas_t = []
for tm in TEM:
    th = temas[tm["tema"]]
    filas_t.append([f"{tm['clave']} · {tm['nombre']}", f"{th['nombre'].split(' ')[0]} {th['version']}", f"{th['rating'] / 20:.1f}".replace(".", ",") + "/5 (" + f"{th['votos']:,}".replace(",", ".") + ")",
                    "Sí" if "Accessibility ready" in th["tags"] else "No", f"{tm['contraste_texto']}:1".replace(".", ","), f"{tm['contraste_banner']}:1".replace(".", ","),
                    f"{tm['contraste_boton']}:1".replace(".", ","), f"{tm['contraste_tarjetas']}:1".replace(".", ","), tm["fidelidad"],
                    "Seleccionada" if tm["clave"] == "A" else "Descartada"])
tabla(["Temática", "Tema candidato", "Valoración", "Accessibility ready", "Contraste texto", "Contraste banner", "Contraste botón", "Contraste tarjetas", "Fidelidad a la identidad del CADUS", "Decisión"],
      filas_t, [2.4, 1.7, 1.7, 1.4, 1.3, 1.3, 1.3, 1.3, 3.3, 1.6], tam=7.5, titulo="Comparativa de las tres temáticas visuales", resaltar=lambda f: f[-1] == "Seleccionada")
rico("**Elección: temática A (institucional), con Blocksy.** Es la única que usa la paleta (granate `#8d0739` y dorado `#faba09`), el logotipo y la tipografía oficiales del CADUS, y todos sus contrastes superan con holgura el mínimo AA (el más bajo es de 7,2:1), con el tema mejor valorado (5,0/5) y con la etiqueta de accesibilidad. La temática B tiene buen contraste pero un aspecto genérico que no reconoce al CADUS (y Astra no declara la etiqueta de accesibilidad). La temática C conserva el dorado pero pierde el granate, y un fondo oscuro es incómodo para leer actas y normativas largas. La comparación de los cinco temas de WordPress con sus datos está en la sección 5.3.")

doc.add_heading("4.5 Evolución del diseño: de los mockups de la entrega 1 al portal final", 2)
par("Para que se vea la evolución del proyecto, esta sección compara los mockups de baja fidelidad de la entrega 1 con lo que se construyó. Primero, la portada en sus tres etapas: el boceto, la primera implementación del Sprint 1 (con el tema Blocksy y un logotipo provisional) y el portal final. Después, cada uno de los cuatro mockups de la entrega 1 (portada, directorio, buzón y ficha de acta) junto a la pantalla implementada, para ver qué se mantuvo y qué cambió al aplicar la identidad del CADUS y los requisitos del Sprint 2.")
figura("evolucion-portada.png", "Evolución de la portada: mockup (entrega 1), primera implementación (Sprint 1) y portal final (Sprint 2). UI-01 · UI-02 · IR-01.", 16.5)
figura("comparativa-1-portada.png", "Portada: mockup de la entrega 1 frente al portal implementado (primer tramo). UI-01 · UI-02 · IR-01.", 16.5)
figura("comparativa-2-directorio.png", "Directorio de delegaciones: mockup frente al portal implementado. UI-03 · RF-06 · RF-07.", 16.5)
figura("comparativa-3-buzon.png", "Buzón de consultas: mockup frente al portal implementado. UI-04 · RF-08 · RF-09 · RN-02 · RN-03.", 16.5)
figura("comparativa-4-acta.png", "Actas y normativa: mockup frente al portal implementado. UI-05 · RF-03 · RF-04 · RF-05.", 16.5)
tabla(["Pantalla", "Mockup de la entrega 1", "Portal implementado", "Qué cambió y por qué", "Requisitos"], [
    ["Portada", "Banner azul genérico con «LOGO», tres comunicados, dos accesos y una franja de próximos eventos.", "Banner granate con el lema y tres contadores reales, presentación de los 40 años, seis accesos, noticias con filtro, eventos, «Nuestro trabajo», galería, buscador y pasos del buzón.", "Se aplicó la paleta, el logotipo y la tipografía oficiales del CADUS y la portada pasó de cuatro bloques a diez, con datos reales del portal.", "UI-01 · UI-02 · IR-01"],
    ["Directorio de delegaciones", "Lista en filas con un desplegable para filtrar por centro y una ficha emergente.", "Tarjetas de color agrupadas por facultad o escuela, cada una con enlace a la ficha de la delegación.", "El agrupamiento por centro sustituye al desplegable y la ficha es una página propia, editable solo por su delegado (RN-04).", "UI-03 · RF-06 · RF-07"],
    ["Buzón de consultas", "Formulario con el recuadro de seguimiento de ticket a su derecha y una nota de privacidad.", "Formulario con consentimiento RGPD y enlace al aviso de privacidad; el seguimiento tiene página propia y la confirmación entrega el código.", "Se separó el seguimiento (exige código y correo, sin caché) y se añadieron el consentimiento, las preguntas frecuentes y el aviso de privacidad.", "UI-04 · RF-08 · RF-09 · RN-02 · RN-03 · RNF-08"],
    ["Actas y normativa", "Ficha con vista previa del PDF, datos del documento y botón de descarga.", "Tabla de documentos con filtro por tipo y botón de descarga; cada documento tiene su ficha.", "El listado con filtro permite encontrar un documento antes de abrir su ficha; el flujo de aprobación (RN-01) se aplica en el panel de administración.", "UI-05 · RF-03 · RF-04 · RF-05 · RN-01"],
], [2.3, 3.7, 4.2, 4.8, 2.0], tam=7.5, titulo="Cambios entre los mockups de la entrega 1 y el portal implementado")
par("Los tres mockups nuevos del Sprint 2 (calendario y votaciones, mediación y panel por rol) no tenían equivalente en la entrega 1: sus pantallas implementadas están en la sección 6.")

# ------------------------------------------------------------------ 5. diseño
doc.add_heading("5. Diseño", 1)
doc.add_heading("5.1 Arquitectura de plugins", 2)
par("La figura muestra qué plugin da respuesta a cada requisito. En verde, los instalados y activos; en naranja, los previstos en el Sprint 1 que finalmente no se instalaron (SupportCandy y YOP Poll), con el motivo en la comparativa.")
figura("uml-diseno-plugins.png", "Diagrama de diseño: plugins de WordPress y requisitos a los que responden.", 16.5)

doc.add_heading("5.2 Comparativa y elección de plugins", 2)
par(f"Para cada tipo de plugin se comparan al menos tres alternativas con los criterios de la guía: fecha de la última actualización, número de instalaciones activas, valoración de los usuarios y número de valoraciones. Los datos proceden de la API pública de wordpress.org consultada el {FECHA_DATOS}; la valoración está en una escala de 5. Además se ha tenido en cuenta la compatibilidad con la versión instalada (WordPress 7.1, PHP 8.x) y el encaje con los requisitos del REM.")
doc.add_heading("Plugins obligatorios: seguridad, caché, analítica y SEO", 3)
comparar("Seguridad", ["wordfence", "all-in-one-wp-security-and-firewall", "better-wp-security", "sucuri-scanner"], "wordfence",
         "Wordfence es el más instalado (más de 5 M), con 4,7/5 en más de 5.000 valoraciones y actualización reciente. Combina cortafuegos, escáner de malware y protección del inicio de sesión, lo que cubre RNF-04 y RNF-08. All-In-One Security (4,7/5) es una alternativa más ligera; Kadence Security (antes Solid Security) y Sucuri están más centrados en endurecimiento y auditoría. La licencia gratuita ya está registrada.")
comparar("Caché", ["wp-fastest-cache", "wp-super-cache", "w3-total-cache", "litespeed-cache"], "wp-fastest-cache",
         "WP Fastest Cache tiene la mejor valoración del grupo (4,9/5 en más de 4.000 valoraciones), se configura con casillas sencillas y funciona sobre Apache con .htaccess, que es nuestro servidor. LiteSpeed Cache es el más instalado, pero su rendimiento depende de un servidor LiteSpeed. W3 Total Cache ofrece más opciones a costa de más complejidad. Responde a RNF-01 (carga inferior a 3 s) y RNF-06.")
comparar("Analítica", ["wp-statistics", "google-site-kit", "google-analytics-for-wordpress", "koko-analytics"], "wp-statistics",
         "WP Statistics guarda los datos en la propia base de datos, sin cookies de terceros ni cuenta de Google, lo que encaja con RNF-08 (RGPD) y funciona desde el servidor local. Site Kit y MonsterInsights tienen más instalaciones pero dependen de Google Analytics y de un sitio público. Koko Analytics es muy ligero pero con menos informes. Para la publicación final se podrá añadir Site Kit.")
comparar("SEO", ["wordpress-seo", "seo-by-rank-math", "all-in-one-seo-pack", "wp-seopress"], "wordpress-seo",
         "Yoast SEO es el estándar del mercado (más de 10 M de instalaciones, 4,8/5 en casi 28.000 valoraciones), con análisis de contenido, metadatos, sitemap XML y migas de pan. Apoya OBJ-02 y UI-02 al mejorar la visibilidad de los comunicados. Rank Math y All in One SEO son alternativas equivalentes; SEOPress es más ligero.")
doc.add_heading("Plugins funcionales", 3)
comparar("Formularios (buzón y mediación)", ["fluentform", "wpforms-lite", "contact-form-7", "ninja-forms"], "fluentform",
         "Fluent Forms guarda las entradas en la base de datos en la versión gratuita (necesario para los tickets del buzón), incluye lógica condicional y validación por expresión regular, con 4,8/5. WPForms Lite solo guarda entradas en la versión de pago y Contact Form 7 solo envía correos. Cubre RF-08, RF-09 y RF-13; la restricción de correo @alum.us.es (RN-02) se aplica en el servidor desde el plugin propio, y se ha comprobado que rechaza un correo ajeno.")
comparar("Tipos de contenido y campos personalizados", ["advanced-custom-fields", "custom-post-type-ui", "pods"], "advanced-custom-fields",
         "ACF aporta los campos de las actas (fecha, pleno y archivo PDF) y de las delegaciones (contacto y horario), con más de 2 M de instalaciones; los campos se definen por código en el plugin propio para que viajen con el repositorio. Los tipos de contenido se registran con la función nativa de WordPress desde ese mismo plugin, de modo que no hace falta Custom Post Type UI, que solo registra tipos y taxonomías. Pods ofrece lo mismo que ACF con una comunidad menor. Cubre RF-03, RF-04, RF-06 y RF-07.")
comparar("Roles y permisos", ["members", "user-role-editor", "wpfront-user-role-editor"], "members",
         "Members tiene la mejor valoración (4,9/5) y permite crear los roles Ejecutiva, Delegado y Estudiante con capacidades a medida, y restringir contenido por rol. Los roles Ejecutiva, Delegado y Estudiante los crea el plugin propio y se gestionan con Members. Son la base de RF-14 y del flujo de aprobación RF-05 / RN-01, que usa el estado nativo «Pendiente de revisión». User Role Editor tiene más instalaciones pero menos funciones de control de contenido.")
comparar("Calendario de eventos", ["the-events-calendar", "my-calendar", "sugar-calendar-lite", "event-organiser"], "the-events-calendar",
         "The Events Calendar es el más instalado (600 k) y tiene vistas de calendario, lista y mes listas para usar; responde a RF-11 y RF-12. Se instaló y configuró en el Sprint 2 (versión 6.18.2), con permisos propios para la Ejecutiva. My Calendar y Event Organiser son alternativas válidas con menos uso.")
comparar("Gestión de tickets", ["supportcandy", "awesome-support", "fluent-support"], "",
         "En la Entrega 1 se seleccionó SupportCandy por su valoración (4,9/5) y su panel por rol. En el Sprint 2 (HU-09) se decidió no instalar ninguno y ampliar el módulo propio del buzón: SupportCandy guarda los tickets en sus propias tablas, así que habría que rehacer la integración con Fluent Forms, el código de seguimiento CAD-AAAA-NNNNN (RF-09) y la página de seguimiento (HU-08), o mantener dos sistemas a la vez. El módulo propio ya cumplía la privacidad de RN-03 (tickets privados, sin API REST ni búsqueda) y solo le faltaban los filtros y el aviso al estudiante, que se han añadido en HU-09a. Así se reduce el riesgo antes de la congelación y no se añade un plugin más que mantener. Awesome Support (3,5/5) y Fluent Support tienen el mismo problema de integración.")
comparar("Votaciones", ["yop-poll", "wp-polls", "poll-maker"], "yop-poll",
         "YOP Poll admite votaciones programadas, control de votos únicos y resultados, necesarios para RF-12 y RN-05; tiene actualización reciente y 4,4/5. En el Sprint 2 no se ha instalado: las convocatorias se publican con apertura y cierre desde el plugin propio y la papeleta con YOP Poll queda como mejora, a confirmar con el cliente.")
comparar("Gestión de documentos descargables", ["download-monitor", "document-library-lite", "wp-document-revisions"], "",
         "Por ahora no se instala ninguno: un campo de archivo de ACF sobre un tipo «Acta» cubre RF-03 y RF-04 sin plugin adicional. Download Monitor (80 k, 4,5/5) se incorporará solo si se necesitan estadísticas de descarga.")

doc.add_heading("5.3 Tema de WordPress", 2)
CT = os.path.join(ENTREGA, "comparativa-temas.md")
for linea in leer_seccion(CT, "Criterios"):
    md_linea(linea)
tt = parsear_tablas(CT)[0]
tabla(["Tema", "Versión", "Última actualización", "Instalaciones activas", "Valoración", "Nº valoraciones", "«Accessibility ready»", "PHP mínimo", "Decisión"],
      [[c.replace("*", "") for c in fila] for fila in tt[1:]], [2.2, 1.5, 2.2, 2.0, 1.7, 1.9, 2.1, 1.7, 1.7], tam=8,
      titulo="Comparativa de temas de WordPress", resaltar=lambda f: f[-1] == "Seleccionado")
par("Se descartó Twenty Twenty-Five (tema por defecto de WordPress) como candidato: solo tiene 14 valoraciones (4,0/5), una base de uso demasiado pequeña para compararlo con el resto. Los datos proceden de la API pública de wordpress.org consultada el 18/10/2026. Los bocetos de las tres temáticas visuales que se compararon están en la sección 4.4.")
for titulo in ("Justificación", "Coste de cambiar de tema", "Limitaciones asumidas"):
    doc.add_heading(titulo, 3)
    for linea in leer_seccion(CT, titulo):
        md_linea(linea)

# ------------------------------------------------------------------ 6. implementación
CAP2 = os.path.join(IMG, "capturas2")
doc.add_heading("6. Implementación", 1)
par("Capturas del portal final (servidor local WampServer, WordPress 7.1, tema Blocksy con el plugin propio CADUS Portal), explicadas una a una. Cada pie indica los requisitos y las historias a las que corresponde. Las páginas públicas se han capturado a página completa sin sesión iniciada; las pantallas de administración, con la cuenta de cada rol. Los textos de noticias, actas y delegaciones son de ejemplo y se sustituirán por contenido real validado con el cliente (sección 8).")
nfig = {}
doc.add_heading("6.1 Portal público", 2)
rico("**Página principal.** Es una sola página con diez bloques que alternan los colores corporativos del CADUS (granate, dorado y gris oscuro): cabecera con el logotipo y el menú principal, banner con el lema «La voz del alumnado» y tres contadores reales (actas y normativas, delegaciones y comunicados publicados), la presentación de los más de 40 años del CADUS, los accesos a las secciones (comunicados, actas, delegaciones, buzón, seguimiento y calendario), las últimas noticias con filtro por categoría, los próximos eventos del calendario, «Nuestro trabajo» (objetivos de la web oficial), la galería de fotos, el buscador, los tres pasos del buzón y el pie con los datos de contacto. Cubre UI-01, UI-02, IR-01, RF-01, RF-11 e HU-01, HU-02 y HU-18.")
trozos = cortes("p01-portada.jpg", CAP2, 5)
for i, caja in enumerate(trozos, 1):
    n_fig = figura("p01-portada.jpg", f"Página principal, parte {i} de {len(trozos)}. UI-01 · UI-02 · IR-01 · IR-02 · HU-01 · HU-02 · HU-18.", 15, CAP2, caja=caja)
    if i == 1:
        nfig["p01-portada.jpg"] = n_fig
    par(["Parte 1: cabecera con logotipo y menú principal, banner con el lema y los tres contadores, y presentación de los más de 40 años del CADUS.",
         "Parte 2: accesos a las seis secciones del portal y últimas noticias con filtro por categoría (Todo, Becas, Eventos, Movilidad y Normativa).",
         "Parte 3: últimos eventos del calendario (de ejemplo) y «Nuestro trabajo» con los objetivos de la web oficial.",
         "Parte 4: galería de fotos, buscador de noticias, actas y delegaciones, y el inicio de «Cómo funciona el buzón».",
         "Parte 5: pasos 2 y 3 del buzón, botón para enviar la consulta y pie con los datos de contacto."][i - 1])
publicas = [
    ("p02-que-es-el-cadus.jpg", "Página «Qué es el CADUS». IR-02 · RNF-07.",
     "Presenta al CADUS y enlaza con su gobierno y estatutos. Lleva un aviso visible de que el texto es provisional hasta que la Delegación General lo valide. Es el segundo elemento del menú principal."),
    ("p03-noticias.jpg", "Noticias y comunicados con filtro por palabra clave, categoría y año. RF-01 · RF-02 · IR-03 · HU-01 · HU-02.",
     "Los comunicados se clasifican en Becas, Movilidad, Normativa y Eventos. El formulario filtra por palabra clave, categoría y año y el listado se pagina. Cada categoría tiene su color."),
    ("p04-actas-y-normativa.jpg", "Repositorio de actas y normativas con filtro por tipo y descarga en PDF. RF-03 · RF-04 · IR-04 · UI-05 · HU-03 · HU-04.",
     "La tabla se ordena por fecha y se filtra por tipo (acta o normativa). Solo se listan los documentos publicados: los que están pendientes de aprobación de la Ejecutiva no aparecen (RN-01). Cada fila tiene su botón de descarga del PDF."),
    ("p05-delegaciones.jpg", "Directorio de delegaciones agrupado por facultad o escuela. RF-06 · RF-07 · IR-05 · UI-03 · HU-05 · HU-06.",
     "Cada centro agrupa sus delegaciones y enlaza con su ficha (delegado, correo, teléfono, horario y enlaces). Hay una ficha de ejemplo por cada uno de los seis centros cargados; faltan las de Química, Física, Biología y la ETSI, que dependen de los datos del CADUS."),
    ("p06-atencion-al-estudiante.jpg", "Atención al estudiante: acceso al buzón, al seguimiento, a la mediación y a las preguntas frecuentes. RF-08 · RF-13 · UI-04.",
     "Reúne en una sola pantalla los cuatro servicios al estudiante con un botón claro en cada tarjeta."),
    ("p07-buzon.jpg", "Buzón de consultas y reclamaciones. RF-08 · RF-09 · IR-06 · RN-02 · RNF-04 · RNF-08 · UI-04 · HU-07.",
     "Formulario con correo universitario, asunto, tipo (consulta o reclamación académica), descripción y consentimiento RGPD obligatorio, con enlace al aviso de privacidad. El servidor comprueba que el correo termina en @alum.us.es (RN-02) y, al enviar, entrega un código de seguimiento con el formato CAD-AAAA-NNNNN seguido de un sufijo aleatorio de 32 caracteres hexadecimales, para que no se pueda adivinar. La reclamación se guarda como contenido privado."),
    ("p08-seguimiento.jpg", "Seguimiento de ticket con código y correo. RF-09 · IR-06 · RN-03 · HU-08.",
     "Exige el código y el correo a la vez. Muestra tipo, asunto, estado y respuestas, sin repetir el texto original. Si los datos no coinciden devuelve un mensaje genérico, sin revelar si el código existe, y limita los intentos. La página no se guarda en caché."),
    ("p09-faq.jpg", "Preguntas frecuentes. RF-08 · RNF-07 · UI-04 · E2-16.",
     "Respuestas desplegables a las dudas habituales: quién puede usar el buzón, diferencia entre consulta y reclamación, cómo saber el estado y qué hacer si se pierde el código."),
    ("p10-privacidad.jpg", "Aviso de privacidad. RNF-08 · RN-03 · UI-04 · E2-16.",
     "Explica el responsable del tratamiento, qué datos se tratan, para qué, la base legal, quién puede verlos, cuánto tiempo se guardan, la seguridad y cómo ejercer los derechos. El plazo de conservación está pendiente de validar con el CADUS."),
    ("p13-calendario.jpg", "Calendario de asambleas y eventos. RF-11 · RF-12 · IR-07 · UI-06 · HU-10.",
     "Calendario de The Events Calendar con vistas de lista, mes y día, buscador y suscripción. Está en el menú principal y alimenta los próximos eventos de la portada. Los eventos de la captura son de ejemplo (llevan «(ejemplo)» en el título); los reales se cargarán con la tarea E2-19."),
    ("p11-votaciones.jpg", "Convocatorias de votación. RF-12 · RN-05 · IR-07 · UI-06 · HU-11.",
     "Lista las convocatorias, cada una vinculada a una asamblea del calendario (RN-05) y con apertura y cierre. La papeleta, que solo se mostraría mientras la votación está abierta, depende de YOP Poll, que no está instalado (sección 5.2): hoy el portal publica las convocatorias, pero no permite emitir el voto, así que RF-12 está parcialmente implementado. La convocatoria de la captura es de ejemplo y está vinculada a la asamblea de ejemplo del calendario."),
    ("p12-mediacion.jpg", "Solicitud de mediación estudiantil. RF-13 · IR-08 · RN-02 · RN-03 · RNF-08 · UI-07 · HU-12.",
     "Formulario con nombre, correo @alum.us.es, facultad o escuela, descripción del conflicto y consentimiento RGPD. Al enviarlo entrega la referencia (MED-AAAA-NNNNN más un sufijo aleatorio de 32 caracteres hexadecimales) sin repetir los datos. La aceptación del tratamiento de datos se conserva con la solicitud. La solicitud es privada: solo la Ejecutiva la ve, no aparece en la web pública, el buscador ni la API, y los avisos por correo no incluyen su contenido."),
    ("p14-contacto.jpg", "Contacto. RNF-07 · UI-02.",
     "Datos del CADUS: dirección en el Pabellón de Uruguay y correo. El mismo bloque de contacto cierra todas las páginas del portal."),
]
for archivo, pie, texto in publicas:
    nfig[archivo] = figura(archivo, pie, 15.5, CAP2)
    par(texto)

doc.add_heading("6.2 Administración y configuración", 2)
admin = [
    (CAP, "20-admin-tickets-filtros.jpg", "Panel de tickets con filtros por estado y tipo. RF-10 · RN-03 · HU-09a.",
     "Solo la Ejecutiva y el administrador ven los tickets. El listado muestra el código de seguimiento en lugar del asunto y permite filtrar por estado y tipo. Los tickets no son públicos ni accesibles por la API REST ni por la búsqueda."),
    (CAP, "21-ticket-aviso-correo.jpg", "Respuesta a un ticket y aviso al estudiante. RF-10 · RN-03 · HU-09a.",
     "Al responder, la Ejecutiva cambia el estado y escribe la respuesta. El estudiante recibe un correo que solo incluye el código y el enlace de seguimiento, nunca el contenido de la respuesta."),
    (CAP, "16-admin-actas.jpg", "Panel de actas con una «Pendiente» de aprobación. RF-05 · RN-01 · HU-03.",
     "El acta pendiente solo se ve aquí y no en la web pública. Solo la Ejecutiva y el administrador gestionan y publican las actas; el delegado no tiene acceso a ellas (RN-01)."),
    (CAP, "19-fluent-forms-buzon.jpg", "Formulario del buzón en Fluent Forms. RF-08 · RN-02 · RNF-04.",
     "El formulario lo crea el plugin propio al activarse e incluye el consentimiento RGPD obligatorio; la validación del dominio y de los campos se repite en el servidor."),
    (CAP, "18-tema-blocksy.jpg", "Pantalla de temas con Blocksy activo. UI-01 · RNF-03 · HU-18.",
     "Blocksy es el tema activo; el aviso sobre su plugin complementario es opcional y no se ha instalado. La identidad visual (paleta, logotipo, menús) se aplica desde el plugin propio."),
    (CAP, "08-yoast-seo-editor.jpg", "Editor de bloques con el panel de Yoast SEO. RF-01 · RF-02 · UI-02.",
     "Yoast SEO se integra en el editor de cada entrada y página para definir frase clave, título y descripción."),
    (CAP, "02-wp-fastest-cache-ajustes.jpg", "Ajustes de WP Fastest Cache. RNF-01 · RNF-06.",
     "Caché de páginas con compresión Gzip; se excluyen los usuarios conectados para que el equipo vea sus cambios al instante y las páginas del buzón no se guardan en caché."),
    (CAP, "06-wordfence.jpg", "Panel de Wordfence con la protección activada. RNF-04 · RNF-08.",
     "Cortafuegos de aplicaciones y escáner de malware con la licencia gratuita registrada."),
    (CAP, "03-wp-statistics.jpg", "Resumen de WP Statistics, analítica sin cookies de terceros. RNF-08 · RNF-01.",
     "Registra visitas y páginas más vistas en la propia base de datos, conforme al RGPD."),
]
admin_nuevas = [
    ("a01-plugins.jpg", "Plugins activos. RF-15 · RNF-01 · RNF-04 · RNF-08.",
     "Plugins obligatorios (Wordfence, WP Fastest Cache, WP Statistics y Yoast SEO), funcionales (ACF, Fluent Forms, Members y The Events Calendar) y el plugin propio CADUS Portal."),
    ("a02-roles.jpg", "Roles del portal en Members. RF-14 · IR-09 · RN-04 · HU-13.",
     "Además de los roles de WordPress existen los del CADUS (miembro de la Ejecutiva, delegado de centro y estudiante), con los permisos descritos en la sección de pruebas."),
    ("a03-escritorio-ejecutivo.jpg", "Escritorio de la Ejecutiva con accesos directos. RF-15 · UI-08 · HU-19b.",
     "Panel de inicio con accesos a noticias, actas, delegaciones, tickets del buzón y calendario, con los avisos de lo pendiente (por ejemplo, un acta que espera aprobación)."),
    ("a04-escritorio-delegado.jpg", "Escritorio y menú del delegado. RF-15 · RN-04 · UI-08 · HU-19a.",
     "El delegado solo ve su ficha de delegación: se ocultan el resto de menús de administración y no tiene acceso a los tickets ni a las mediaciones. La cuenta de prueba de la captura no tenía ninguna delegación asignada, por eso el panel lo indica."),
    ("a06-eventos.jpg", "Formulario de nuevo evento del calendario. RF-11 · HU-10.",
     "La Ejecutiva crea asambleas y eventos con título, descripción, fecha y hora de inicio y fin, recinto y organizadores. El menú «Eventos» solo aparece a los roles con permiso, y lo creado se muestra en el calendario y en la portada."),
    ("a07-votaciones.jpg", "Formulario de convocatoria de votación. RF-12 · RN-05 · HU-11.",
     "Cada convocatoria exige elegir una asamblea del calendario (RN-05) y fijar la apertura y el cierre; sin ambos datos no se puede publicar y se guarda como borrador. El identificador de la encuesta de YOP Poll es opcional."),
]
for carpeta, archivo, pie, texto in admin:
    nfig[archivo] = figura(archivo, pie, 15.5, carpeta)
    par(texto)
for archivo, pie, texto in admin_nuevas:
    if os.path.exists(os.path.join(CAP2, archivo)):
        nfig[archivo] = figura(archivo, pie, 15.5, CAP2)
        par(texto)

doc.add_heading("Estado de la implementación por historia de usuario", 2)
F = lambda *a: ", ".join(f"figura {nfig[x]}" for x in a if x in nfig) or "—"
tabla(["Historia", "Estado", "Evidencia"], [
    ["PB-002 · PB-005 · PB-006 · PB-007 · PB-008", "Hecho", "Servidor Wamp con acceso por Tailscale, servidor de Discord, plugins y REM."],
    ["HU-01 · HU-02", "Hecho", f"Portada y noticias con filtros ({F('p01-portada.jpg', 'p03-noticias.jpg')})."],
    ["HU-03 · HU-04", "Hecho", f"Flujo de aprobación y repositorio con descarga ({F('16-admin-actas.jpg', 'p04-actas-y-normativa.jpg')})."],
    ["HU-05 · HU-06", "Hecho", f"Directorio y ficha de delegación ({F('p05-delegaciones.jpg')}); edición restringida al delegado propietario."],
    ["HU-07 · HU-08", "Hecho", f"Buzón, código de ticket y seguimiento ({F('p07-buzon.jpg', 'p08-seguimiento.jpg')})."],
    ["HU-09 (9a y 9b)", "Hecho", f"Panel de tickets con filtros, aviso al estudiante y privacidad RN-03 ({F('20-admin-tickets-filtros.jpg', '21-ticket-aviso-correo.jpg')})."],
    ["HU-10", "Hecho", f"Calendario de asambleas y eventos ({F('p13-calendario.jpg')})."],
    ["HU-11", "Hecho (sin papeleta)", f"Convocatorias de votación vinculadas a una asamblea ({F('p11-votaciones.jpg')}); falta la papeleta para votar (YOP Poll no instalado), que pasa a HU-21 en el Sprint 3."],
    ["HU-12", "Hecho", f"Solicitud de mediación privada ({F('p12-mediacion.jpg')})."],
    ["HU-13", "Hecho", f"Roles de WordPress y del CADUS ({F('a02-roles.jpg')})."],
    ["HU-18", "Hecho", f"Tema Blocksy con la identidad del CADUS ({F('p01-portada.jpg', '18-tema-blocksy.jpg')})."],
    ["HU-19 (19a, 19b y 19c)", "Hecho", f"Menús por rol, escritorio con accesos y pruebas por rol ({F('a03-escritorio-ejecutivo.jpg', 'a04-escritorio-delegado.jpg')}); la documentación de las pruebas de HU-19c y HU-12 está en la carpeta tests/ (PR #32 y #33, fusionadas)."],
    ["HU-14 a HU-17 · HU-20 · HU-21", "Pendiente", "Sprints 3 y 4 (móvil, accesibilidad WCAG, filtros del panel, protección frente a spam, rendimiento y papeleta de votación)."],
], [4.4, 2.0, 10.6], titulo="Estado de la implementación")

doc.add_heading("Pruebas realizadas", 2)
par("Las pruebas se hicieron en el servidor compartido con datos ficticios (correos de prueba, cuentas temporales) que se eliminan al terminar; cada batería comprueba al final que no queda ningún dato de prueba.")
tabla(["Prueba", "Resultado esperado", "Resultado"], [
    ["Enviar el buzón con un correo de gmail.com", "Rechazo con mensaje (RN-02)", "Correcto: «Utiliza un correo universitario válido terminado en @alum.us.es»"],
    ["Enviar el buzón y la mediación con un correo @alum.us.es", "Código CAD-AAAA-NNNNN o referencia MED-AAAA-NNNNN, ambos con sufijo aleatorio (RF-09, RF-13)", "Correcto"],
    ["Seguimiento con código correcto y correo equivocado", "Sin datos del ticket (RN-03)", "Correcto: mensaje genérico, no revela si el código existe"],
    ["Batería de privacidad de tickets (privacidad_tickets.py)", "Ningún dato del ticket sale en URLs, búsqueda, feeds, mapas del sitio ni API REST para anónimos ni para ningún rol salvo Ejecutiva y administrador", "40 comprobaciones correctas, 0 fallos"],
    ["Batería HTTP RN-03 por rol (rn03_http.py) y de mediaciones (rn03_mediacion.py)", "Ningún rol sin permiso ve tickets ni mediaciones, ni sus adjuntos por la API REST o AJAX; un error 500 no cuenta como correcto", "0 fallos y sin datos de prueba al terminar"],
    ["Permisos de cada rol sobre actas, delegaciones, tickets y calendario", "Solo Ejecutiva y administrador publican y ven tickets; el delegado edita solo su ficha (RN-01, RN-03, RN-04)", "Correcto en todos los roles"],
    ["Páginas públicas (interiores_http.py)", "Un único título h1 y un único pie por página", "34 páginas HTML inspeccionadas: todas con un único pie; 32 con un único h1. La portada y el calendario tienen un segundo h1 (el del tema y el de The Events Calendar); se corrige en HU-15 (accesibilidad)"],
    ["Volver a ejecutar la configuración del plugin", "No duplica contenido ni quita permisos", "Correcto (los permisos del calendario se reponen solos si se pierden)"],
], [5.5, 6.0, 5.5], tam=8, titulo="Pruebas funcionales y de seguridad")

# ------------------------------------------------------------------ 7. trazabilidad
doc.add_heading("7. Matriz de trazabilidad", 1)
par("Cada requisito del REM aparece con el mismo código en el análisis (entidad o pantalla), en el diseño (plugin) y en la implementación (figura). Los números de figura se recalculan al regenerar el documento.")
tabla(["Requisito (REM)", "Análisis (UML / mockup)", "Diseño (plugin)", "Implementación"], [
    ["RF-01 Publicación de noticias", "Clases Noticia, Categoria · UI-02", "Plugin propio + Yoast SEO", F("p03-noticias.jpg", "08-yoast-seo-editor.jpg")],
    ["RF-02 Búsqueda de noticias", "Clase Noticia · UI-02", "Plugin propio (filtro)", F("p03-noticias.jpg")],
    ["RF-03 Repositorio de actas", "Clase DocumentoInstitucional · UI-05", "Plugin propio + ACF", F("p04-actas-y-normativa.jpg")],
    ["RF-04 Repositorio de normativas", "Clase DocumentoInstitucional · UI-05", "Plugin propio + ACF", F("p04-actas-y-normativa.jpg")],
    ["RF-05 Aprobación previa", "DocumentoInstitucional.estadoAprobacion", "Plugin propio + Members", F("16-admin-actas.jpg")],
    ["RF-06 Directorio de delegaciones", "Clases Delegacion, FacultadEscuela · UI-03", "Plugin propio + ACF", F("p05-delegaciones.jpg")],
    ["RF-07 Ficha de delegación", "Clase Delegacion · UI-03", "Plugin propio + ACF + Members", F("p05-delegaciones.jpg")],
    ["RF-08 Buzón de consultas", "Clase Ticket · UI-04", "Fluent Forms + plugin propio", F("p07-buzon.jpg", "19-fluent-forms-buzon.jpg")],
    ["RF-09 Código de seguimiento", "Ticket.codigoSeguimiento · UI-04", "Plugin propio", F("p07-buzon.jpg", "p08-seguimiento.jpg")],
    ["RF-10 Gestión de tickets", "Clases Ticket, RespuestaTicket", "Plugin propio (módulo buzón)", F("20-admin-tickets-filtros.jpg", "21-ticket-aviso-correo.jpg")],
    ["RF-11 Calendario", "Clase EventoCalendario · UI-06", "The Events Calendar", F("p13-calendario.jpg", "a06-eventos.jpg")],
    ["RF-12 Convocatorias de votación", "EventoCalendario.asambleaVinculada · UI-06", "Plugin propio (YOP Poll opcional)", F("p11-votaciones.jpg", "a07-votaciones.jpg") + " (parcial: sin papeleta)"],
    ["RF-13 Solicitud de mediación", "Clase SolicitudMediacion · UI-07", "Fluent Forms + plugin propio", F("p12-mediacion.jpg", "a05-mediaciones.jpg")],
    ["RF-14 Roles y permisos", "Clase Usuario (rol) · UI-08", "Plugin propio + Members", F("a02-roles.jpg")],
    ["RF-15 Panel de administración", "Todas las clases · UI-08", "Panel de WordPress + ACF + Members", F("a01-plugins.jpg", "a03-escritorio-ejecutivo.jpg", "a04-escritorio-delegado.jpg")],
    ["RNF-01 Rendimiento", "Requisito de calidad: sin elemento propio en el análisis", "WP Fastest Cache", F("02-wp-fastest-cache-ajustes.jpg") + "; medición del tiempo de carga pendiente (HU-20)"],
    ["RNF-02 Usabilidad móvil", "Mockups (diseño adaptable)", "Tema Blocksy + CSS propio", "Capturas a 1366 px; la comprobación en móvil y tablet está pendiente (HU-14, Sprint 3)"],
    ["RNF-03 Accesibilidad", "Mockups", "Tema Blocksy (declarado «Accessibility ready»)", F("18-tema-blocksy.jpg") + "; la verificación WCAG 2.1 AA está pendiente (HU-15, Sprint 3)"],
    ["RNF-04 Seguridad en formularios", "Requisito de calidad: sin elemento propio en el análisis", "Wordfence + Fluent Forms + validación en el servidor", F("06-wordfence.jpg", "p07-buzon.jpg") + "; baterías de pruebas de la sección 6"],
    ["RNF-05 Compatibilidad técnica", "Restricción del entorno RC-01", "WampServer: Apache, PHP 8.x y MySQL 8.x", "WordPress 7.1 funcionando sobre ese servidor (todas las capturas)"],
    ["RNF-06 Disponibilidad", "Restricción del entorno RC-01", "WP Fastest Cache + servidor compartido por Tailscale", F("02-wp-fastest-cache-ajustes.jpg") + "; limitada a que el equipo anfitrión esté encendido (sección 2.2.1)"],
    ["RNF-07 Mantenibilidad", "Restricciones RC-02 y RC-03", "Plugin propio documentado, scripts y repositorio GitHub", "README del plugin, carpeta scripts/ y docs/ del repositorio (sin captura)"],
    ["RNF-08 Privacidad", "Ticket.confidencial · UI-04 · UI-07", "WP Statistics + Wordfence + plugin propio", F("03-wp-statistics.jpg", "p10-privacidad.jpg", "p08-seguimiento.jpg") + "; pendientes el plazo de conservación y las copias de seguridad"],
    ["UI-01 Diseño gráfico", "Mockup de portada", "Tema Blocksy + plugin propio", F("p01-portada.jpg", "18-tema-blocksy.jpg")],
    ["UI-02 Página principal", "Mockup 1 · estructura web", "Tema + Yoast SEO", F("p01-portada.jpg")],
    ["UI-03 Directorio", "Mockup 2", "Plugin propio", F("p05-delegaciones.jpg")],
    ["UI-04 Buzón de consultas", "Mockup 3", "Fluent Forms + plugin propio", F("p07-buzon.jpg", "p09-faq.jpg", "p10-privacidad.jpg")],
    ["UI-05 Ficha de acta o normativa", "Mockup 4", "Plugin propio + ACF", F("p04-actas-y-normativa.jpg")],
    ["UI-06 Calendario y votaciones", "Mockup 5", "The Events Calendar + plugin propio", F("p13-calendario.jpg", "p11-votaciones.jpg")],
    ["UI-07 Solicitud de mediación", "Mockup 6", "Fluent Forms + plugin propio", F("p12-mediacion.jpg")],
    ["UI-08 Panel de administración por rol", "Mockup 7", "Plugin propio + Members", F("a03-escritorio-ejecutivo.jpg", "a04-escritorio-delegado.jpg")],
    ["IR-01 Página principal", "Mockup 1 · estructura web (Inicio)", "Tema Blocksy + plugin propio (inicio)", F("p01-portada.jpg")],
    ["IR-02 Menús", "Estructura del sitio web (sección 4.1)", "Menús de WordPress configurados por el plugin propio", F("p01-portada.jpg")],
    ["IR-03 Noticias y comunicados", "Clases Noticia y Categoria", "Plugin propio + Yoast SEO", F("p03-noticias.jpg")],
    ["IR-04 Actas y normativas", "Clase DocumentoInstitucional", "Plugin propio + ACF", F("p04-actas-y-normativa.jpg", "16-admin-actas.jpg")],
    ["IR-05 Delegaciones de alumnos", "Clases Delegacion y FacultadEscuela", "Plugin propio + ACF", F("p05-delegaciones.jpg")],
    ["IR-06 Consultas y reclamaciones", "Clases Ticket y RespuestaTicket", "Fluent Forms + plugin propio (buzón)", F("p07-buzon.jpg", "p08-seguimiento.jpg", "20-admin-tickets-filtros.jpg")],
    ["IR-07 Calendario", "Clase EventoCalendario", "The Events Calendar + plugin propio (votaciones)", F("p13-calendario.jpg", "p11-votaciones.jpg", "a06-eventos.jpg")],
    ["IR-08 Solicitudes de mediación", "Clase SolicitudMediacion", "Fluent Forms + plugin propio (mediación)", F("p12-mediacion.jpg")],
    ["IR-09 Usuarios y roles", "Clase Usuario", "Members + plugin propio (roles)", F("a02-roles.jpg")],
    ["RN-01 Aprobación de documentos", "DocumentoInstitucional.estadoAprobacion", "Plugin propio + Members (solo la Ejecutiva publica)", F("16-admin-actas.jpg")],
    ["RN-02 Correo institucional", "Ticket y SolicitudMediacion (correo @alum.us.es)", "Fluent Forms + validación en el servidor del plugin propio", F("p07-buzon.jpg", "p12-mediacion.jpg") + "; prueba de rechazo en la sección 6"],
    ["RN-03 Privacidad de reclamaciones", "Ticket.confidencial", "Plugin propio (contenido privado, sin API ni búsqueda)", F("p08-seguimiento.jpg", "20-admin-tickets-filtros.jpg") + "; baterías de pruebas de la sección 6"],
    ["RN-04 Edición de fichas", "Delegacion (autor = delegado)", "Plugin propio + Members (rol de delegado)", F("a04-escritorio-delegado.jpg", "p05-delegaciones.jpg")],
    ["RN-05 Votaciones vinculadas", "EventoCalendario.asambleaVinculada", "Plugin propio (votaciones) + The Events Calendar (categoría Asambleas)", F("a07-votaciones.jpg", "p11-votaciones.jpg")],
], [4.3, 4.7, 4.3, 3.7], tam=8, titulo="Trazabilidad requisito → análisis → diseño → implementación")

# ------------------------------------------------------------------ 8. calidad e incidencias
doc.add_heading("8. Calidad, incidencias y pendientes", 1)
doc.add_heading("Revisión cruzada del código", 2)
par("El plugin propio se reparte entre dos asistentes de programación: Claude escribe la mayor parte y Codex escribe módulos y revisa. Después cada uno revisó el trabajo del otro y el resultado se probó en el servidor; el código se revisó y comprendió antes de integrarlo (desglose en el README del plugin).")
tabla(["Revisión", "Hallazgos", "Resolución"], [
    ["Sprint 1. Codex sobre el código de Claude", "9: error 500 con ?tipo[]=; borrado de contenido por defecto en cada activación; ajustes sobrescritos al reactivar; ejemplos recreados al renombrarlos; la búsqueda incluía páginas y actas; actas sin fecha omitidas; enlace al PDF perdido sin ACF; orden de enlaces; roles borrados y recreados", "Los 9 corregidos y verificados"],
    ["Sprint 1. Claude sobre el código de Codex", "Sin fallos bloqueantes. El código del ticket no es correlativo y el seguimiento no acredita la propiedad del correo", "Aceptado y documentado; se mitigó con límite de intentos y exigiendo código y correo"],
    ["Sprint 2. Codex sobre el diseño y la portada", "Contrastes ilegibles en el texto de las tarjetas al pasar el ratón y el autor «root» visible en las fichas", "Corregidos; el marcado de las páginas públicas se inspecciona con una batería (un solo pie en todas; dos páginas con un h1 de más, ver pruebas)"],
    ["Sprint 2. Codex sobre toda la entrega (25/10)", "8: adjuntos privados accesibles por AJAX (alto); el importador de cuentas podía modificar otra cuenta (alto); autorreparación del calendario incompleta; pruebas que daban por bueno un error 500; reapertura de tareas en el script de Scrum; foco invisible en tarjetas oscuras; el Project no recuperaba elementos incompletos; resumen que afirmaba una integración pendiente", "Los 8 corregidos y las baterías de pruebas, repetidas sin fallos"],
], [4.0, 8.6, 4.4], tam=8, titulo="Revisión cruzada")
doc.add_heading("Incidencias resueltas durante los sprints", 2)
tabla(["Incidencia", "Causa", "Solución"], [
    ["WordPress no podía descargar plugins (cURL error 77)", "Faltaba wp-includes/certificates/ca-bundle.crt: el .gitignore excluye *.crt", "Restaurado el archivo desde el repositorio oficial de WordPress"],
    ["Pantallas de Yoast y Fluent Forms en blanco", "Faltaban wp-includes/js/dist, css/dist y build: el .gitignore excluye dist/ y build/", "Restauradas desde el paquete oficial de WordPress 7.1"],
    ["WP Fastest Cache no guardaba sus ajustes", "No existía el archivo .htaccess", "Creado el .htaccess estándar de WordPress"],
    ["La Ejecutiva perdía los permisos de eventos", "Otro proceso reescribió la lista de roles con una copia antigua", "El módulo del calendario comprueba y repone todos los permisos en cada carga"],
    ["Una rama de mediación habría eliminado la portada, los menús por rol y el diseño", "Se hizo sobre una copia antigua del plugin", "Se integró solo su módulo y los tres cambios de configuración necesarios"],
    ["La vista previa de impresión de REMUS salía en blanco", "Sin confirmar: el HTML que genera REMUS es correcto y se ve bien en un navegador, así que falla el visor interno de REMUS", "El PDF del REM se obtiene imprimiendo ese HTML desde Chrome"],
], [4.8, 6.0, 6.2], tam=8, titulo="Incidencias técnicas y su resolución")
doc.add_heading("Pendientes y decisiones abiertas", 2)
tabla(["Pendiente", "Motivo", "Responsable"], [
    ["Entrevista con el CADUS y validación del REM", "Depende de la disponibilidad del cliente; condiciona el plazo de conservación de datos, la lista de cuentas y la identidad visual", "Miguel Ángel"],
    ["Contenido real (noticias, actas, estatutos, delegaciones)", "El CADUS no lo ha entregado; todo es de ejemplo y está marcado como tal", "Miguel Ángel · Joaquín"],
    ["Fichas de Química, Física, Biología y ETSI", "Faltan los datos de contacto de esas delegaciones", "Álvaro"],
    ["Mockups de alta fidelidad en Figma", "Los siete mockups y las tres temáticas están en Figma como bocetos de baja fidelidad; falta un diseño de alta fidelidad con componentes, que no se ha hecho", "Miguel Ángel"],
    ["Cuentas de los delegados (una por delegado, con su correo de alumno)", "La cuenta única del CADUS (dcadus@us.es) ya está creada; las de los delegados se crean con el script cuando el CADUS facilite la lista", "Pablo"],
    ["Papeleta de las votaciones (YOP Poll) y eventos reales en el calendario", "No instalado; las convocatorias se publican sin papeleta. La emisión del voto es la historia HU-21 (Sprint 3), pendiente de validar con el cliente", "Joaquín · Miguel Ángel"],
    ["Plazo de conservación y copias de seguridad de tickets y mediaciones (RNF-08)", "Pendiente de validar con el CADUS", "Pablo"],
    ["Volcado de la base de datos del Sprint 2 y revisión de Wordfence", "Se hace al cerrar la entrega", "Pablo"],
    ["Sprint 3: HU-14 a HU-17 (móvil, accesibilidad WCAG 2.1 AA, filtros del panel y protección frente a spam)", "Planificado del 26/10 al 15/11 en el GitHub Project", "Equipo"],
], [6.0, 8.0, 3.0], tam=8, titulo="Pendientes y decisiones abiertas")

# ------------------------------------------------------------------ pie de página
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Portal CADUS · Entrega 2 (CMS) · Página ")
r.font.size = Pt(8.5)
campo(fp, "PAGE", "1")

# actualizar campos (índice) al abrir
settings = doc.settings.element
u = OxmlElement("w:updateFields")
u.set(qn("w:val"), "true")
settings.append(u)

doc.core_properties.title = "Portal web del CADUS — Entrega 2 (CMS, versión final)"
doc.core_properties.author = "Equipo GSI — CADUS"
doc.save(SALIDA)
print("documento generado:", SALIDA, "·", figura_n[0], "figuras ·", tabla_n[0], "tablas")
